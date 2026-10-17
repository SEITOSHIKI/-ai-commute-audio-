"""Rewrite the per-slide sections of a 台本 md from the deck's speaker notes (the header part of the md is kept).
Usage: python3 tools/notes_md.py deck.pptx script.md
Each note is "【開始 - 終了】\\n要約：見出し\\n\\n本文"; the md gets "## n. 見出し　【開始 - 終了】\\n\\n本文"."""
import re, sys
from pptx import Presentation

pptx_path, md_path = sys.argv[1], sys.argv[2]
head = open(md_path, encoding="utf-8").read()
m = re.search(r"^## 1\. ", head, re.M)
head = head[: m.start()] if m else head.rstrip() + "\n\n"
out = []
for i, slide in enumerate(Presentation(pptx_path).slides, 1):
    t = slide.notes_slide.notes_text_frame.text.strip() if slide.has_notes_slide else ""
    lines = t.split("\n")
    when = lines[0] if lines and lines[0].startswith("【") else ""
    title = next((l[3:] for l in lines if l.startswith("要約：")), f"スライド{i}")
    body = "\n".join(l for l in lines if l.strip() and l != when and not l.startswith("要約：")).strip()
    out.append(f"## {i}. {title}　{when}\n\n{body}\n")
open(md_path, "w", encoding="utf-8").write(head + "\n".join(out))
print(md_path, len(out), "slides")
