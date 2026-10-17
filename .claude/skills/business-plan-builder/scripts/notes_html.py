"""Build a printable HTML (A4 portrait): each slide image on top, its speaker notes (台本) below.
Usage: python3 notes_html.py deck.pptx slides_dir out.html "表示名"
slides_dir holds s-01.png, s-02.png ... rendered from the deck's PDF."""
import html, sys, pathlib
from pptx import Presentation

pptx_path, img_dir, out, title = sys.argv[1], pathlib.Path(sys.argv[2]), sys.argv[3], sys.argv[4]
prs = Presentation(pptx_path)
imgs = sorted(img_dir.glob("s-*.png"))
pages = []
for i, (slide, img) in enumerate(zip(prs.slides, imgs), 1):
    notes = slide.notes_slide.notes_text_frame.text if slide.has_notes_slide else ""
    lines = notes.strip().split("\n")
    head, msg, body = "", "", []
    for ln in lines:
        if ln.startswith("【") and not head:
            head = ln
        elif ln.startswith("要約：") and not msg:
            msg = ln[3:]
        elif ln.strip():
            body.append(ln)
    body_html = "".join(f"<p>{html.escape(b)}</p>" for b in body) or "<p class='none'>（台本なし）</p>"
    pages.append(f"""<section class="page">
  <div class="top"><span>{html.escape(title)}</span><span>{i} / {len(prs.slides)}</span></div>
  <img src="{img.resolve().as_uri()}" alt="スライド{i}">
  <div class="meta">{f'<span class="time">{html.escape(head)}</span>' if head else ''}{f'<span class="msg">{html.escape(msg)}</span>' if msg else ''}</div>
  <div class="notes">{body_html}</div>
</section>""")

css = """
@page { size: A4 portrait; margin: 12mm 14mm; }
* { box-sizing: border-box; }
body { margin: 0; font-family: 'IPAPGothic','IPAGothic',sans-serif; color: #1a1a1a; }
.page { break-after: page; }
.page:last-child { break-after: auto; }
.top { display: flex; justify-content: space-between; font-size: 8.5pt; color: #808080; border-bottom: 0.6pt solid #bfbfbf; padding-bottom: 2mm; margin-bottom: 3mm; }
img { width: 100%; border: 0.6pt solid #bfbfbf; display: block; }
.meta { margin: 4mm 0 2mm; display: flex; flex-direction: column; gap: 1.2mm; }
.time { font-size: 9pt; color: #b8692e; font-weight: bold; }
.msg { font-size: 11pt; font-weight: bold; border-left: 2.4pt solid #b8692e; padding-left: 2.5mm; }
.notes p { font-size: 10.5pt; line-height: 1.75; margin: 0 0 2mm; text-align: justify; }
.none { color: #808080; }
"""
pathlib.Path(out).write_text(f"<!doctype html><html lang='ja'><head><meta charset='utf-8'><style>{css}</style></head><body>{''.join(pages)}</body></html>", encoding="utf-8")
print("slides", len(prs.slides), "images", len(imgs))
