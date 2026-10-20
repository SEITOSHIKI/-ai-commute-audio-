import json, subprocess, sys, shutil
from openpyxl import load_workbook
W = sys.argv[1]
REG = json.load(open(W + "/reg.json", encoding="utf-8")); NV = json.load(open(W + "/node_values.json", encoding="utf-8")); IN = json.load(open(W + "/inputs.json", encoding="utf-8"))
import os
SCRIPTS = os.path.join(os.environ["SK"], "xlsx", "scripts")  # make_all.sh が渡す
RECALC = os.path.join(SCRIPTS, "recalc.py")
G, EXT = IN["GROWTH"], IN["EXT"]
ns = G["newSites"] + EXT["newSites"]; ct = G["contract"] + EXT["contract"]
def setp(ws, key, vals):
    sh, r, c = REG[key]
    if not isinstance(vals, list): vals = [vals]
    for i, v in enumerate(vals): ws.cell(r, c + i).value = v
k = 0.7
SC = [
  {},
  {"k_price": 0.8},
  {"k_churn": 0.2},
  {"k_cost": 2},
  {"k_salary": 100},
  {"new_sites": [round(n * k + 1e-9) for n in ns], "contract": [c * k for c in ct], "k_hire": k},
  {"new_sites": [round(n * k + 1e-9) for n in ns], "contract": [c * k for c in ct]},
  {"devices": 15, "glass_n": 12, "disp_n": 3},
  {"k_cost": 60 / 7.199999999999999},
  {"head5": [0, 0, 1, 3, 3] + EXT["heads"][5]},
]
res = NV["SC"] + [{"n": "営業の採用を3年目から", **NV["salesY3"]}]
def val(wb, key, y=1):
    sh, r, c = REG[key]; return wb[sh].cell(r, c + y - 1).value
allok = True
for i, sc in enumerate(SC):
    wb = load_workbook(W + "/test.xlsx")
    ws = wb["前提"]
    for kk, vv in sc.items(): setp(ws, kk, vv)
    f = f"{W}/scen/s{i}.xlsx"; wb.save(f)
    for attempt in range(3):
        out = subprocess.run(["python3", RECALC, f, "120"], capture_output=True, text=True, cwd=SCRIPTS)
        j = json.loads(out.stdout)
        if "status" in j: break
        print("  retry:", j)
        wb.save(f)
    st = j.get("status")
    wv = load_workbook(f, data_only=True)
    o5, o10 = res[i]["o5"], res[i]["o10"]
    got = {"fb5": val(wv, "first_black5"), "op5": val(wv, "op", 5), "min5": val(wv, "min_cash5"), "fb10": val(wv, "first_black10"), "cb10": val(wv, "cum_back10"), "min10": val(wv, "min_cash10")}
    exp = {"fb5": o5["firstBlack"] or "5年でなし", "op5": o5["op5"], "min5": o5["minCash"], "fb10": o10["firstBlack"] or "10年でなし", "cb10": o10["cumBack"] or "10年でなし", "min10": o10["minCash"]}
    ok = all((abs(got[x] - exp[x]) < 1e-6 if isinstance(exp[x], (int, float)) and isinstance(got[x], (int, float)) else got[x] == exp[x]) for x in exp)
    # per-year op and cash
    ok2 = all(abs(val(wv, "op", y) - o10["ops"][y - 1]) < 1e-6 and abs(val(wv, "cash", y) - o10["cash"][y - 1]) < 1e-6 for y in range(1, 11))
    allok &= ok and ok2
    print(i, res[i]["n"], st, "OK" if ok and ok2 else "NG", {x: (round(got[x], 3) if isinstance(got[x], float) else got[x]) for x in got}, "" if ok else exp)
print("ALL OK" if allok else "SOME NG")
