# 再計算したブックの値を読み、tora_v2.js の値（node_values.json）と突き合わせる
import json, sys
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter as CL
WORK, X = sys.argv[1], sys.argv[2]
REG = json.load(open(WORK + "/reg.json", encoding="utf-8"))
NV = json.load(open(WORK + "/node_values.json", encoding="utf-8"))
wb = load_workbook(X, data_only=True)
def v(key, y=1):
    sh, r, c = REG[key]
    return wb[sh].cell(r, c + y - 1).value
bad = []; n = 0
def chk(label, a, b, tol=1e-6):
    global n
    n += 1
    if isinstance(b, (int, float)) and isinstance(a, (int, float)):
        if abs(a - b) > tol * max(1, abs(b)): bad.append((label, a, b))
    elif a != b: bad.append((label, a, b))
R = NV["run10"]["rows"]
M = {"open": "open", "churn_n": "churn", "new": "nw", "close": "sites", "lic_exist": "existing", "lic_new": "newLic", "license": "license",
     "setup_rev": "setup", "contract_rev": "contract", "service": "service", "sales": "sales", "lic_cost": "licCost", "setup_cost_y": "setupCost",
     "setup_labor_y": "setupLabor", "contract_labor": "contractLabor", "cogs": "cogs", "gp": "gp", "gp_rate": "gpRate", "payroll_pl": "payroll",
     "recruit_pl": "recruit", "transfer": "transfer", "sga": "sga", "op": "op", "cum": "cum", "cash": "cash", "staff": "staff",
     "bep": "bep", "sites_per": "sitesPerStaff", "lic_margin": "licMargin", "funding_y": "funding", "sga_per": "perStaff"}
for k, nk in M.items():
    for y in range(1, 11):
        chk(f"{k} y{y}", v(k, y), R[y - 1][nk])
# other は tora_v2.js では顧問込み
for y in range(1, 11):
    chk(f"other+adv y{y}", v("other", y) + v("adv", y), R[y - 1]["other"])
T = NV["team"]
for y in range(1, 11):
    chk(f"staff_plan y{y}", v("staff_plan", y), T[y-1]["staff"]); chk(f"pay_total y{y}", v("pay_total", y), T[y-1]["payroll"]); chk(f"recruit_plan y{y}", v("recruit_plan", y), T[y-1]["recruit"])
    for i in range(8):
        chk(f"sal{i} y{y}", v(f"sal{i}_y", y), T[y-1]["rows"][i]["salary"]); chk(f"h{i} y{y}", v(f"h{i}", y), T[y-1]["rows"][i]["heads"])
cv = NV["cv"]
for k, nk in [("per_day_man", "perDeviceDay"), ("docs", "docs"), ("rw", "rework"), ("cost", "cost"), ("ratio", "ratio"), ("ratio_up", "ratioUp"), ("be", "breakEvenMin"), ("freed_h", "freedHours")]:
    chk(k, v(k), cv[nk])
chk("hW", v("j_h", 1), NV["H"]["electrician"] * 10000); chk("hM", v("j_h", 2), NV["H"]["manager"] * 10000)
# sensitivity (days x minutes)
for di, row in enumerate(NV["sensDays"], 1):
    for mi, cell in enumerate(row, 1):
        chk(f"sens eff d{di} m{mi}", v(f"sens_eff{di}", mi), cell["docs"]); chk(f"sens ratio d{di} m{mi}", v(f"sens_ratio{di}", mi), cell["ratio"])
    chk(f"sens be d{di}", v(f"sens_be{di}", 1), row[0]["be"])
for i, s in enumerate(NV["sens"], 1):
    pass
U = NV["unit"]
chk("gp_site", v("gp_site"), U["gpYear"]); chk("setup_gp", v("setup_gp"), U["setupGp"]); chk("cac", v("cac"), U["cac"]); chk("payback", v("payback"), U["paybackMonths"])
chk("lc_pay/3", v("lc_pay") / 3, U["salesPay"])
for i, y in enumerate([3, 5, 7, 10], 1): chk(f"ltv{y}", v(f"ltv_t{i}", 1), NV["ltv"][i-1])
chk("cogs_site", v("cogs_site"), NV["cogsSite"]); chk("margin", v("margin_site"), NV["marginSite"]); chk("perdev", v("perdev"), NV["cloudPerDevice"])
chk("csday", v("csday"), NV["CSDAY"]); chk("setup_cost", v("setup_cost"), NV["setupCost"]); chk("lease", v("pay_lease"), NV["leaseSite"])
for i, m in enumerate(NV["margins"], 1): chk(f"mg{i} part", v(f"mg{i}", 2), m["part"]); chk(f"mg{i} full", v(f"mg{i}", 3), m["full"])
chk("first_black10", v("first_black10"), NV["run10"]["firstBlack"]); chk("cum_back10", v("cum_back10"), NV["run10"]["cumBack"]); chk("min_cash10", v("min_cash10"), NV["run10"]["minCash"])
chk("min_cash5", v("min_cash5"), NV["run5"]["minCash"]); chk("first_black5", v("first_black5"), "5年でなし"); chk("cum_bottom", v("cum_bottom"), NV["run10"]["cumPeak"])
mk = NV["market"]
chk("sam", v("sam"), mk["sam"]); chk("tam", v("tam"), mk["tamSeg"]); chk("sam_n", v("sam_n"), mk["samSites"]); chk("som_sites", v("som_sites"), mk["somSites"]); chk("som_value", v("som_value"), mk["somValue"])
for i, c in enumerate(mk["cells"]): chk(f"seg{i} value", wb["市場規模"].cell(REG[f"mk{i}"][1], 6).value, c["value"])
print("checked", n, "mismatches", len(bad))
for b in bad[:40]: print("  MISMATCH", b)
# 照合シート
ws = wb["照合"]
res = [ws.cell(r, 8).value for r in range(5, ws.max_row + 1) if ws.cell(r, 8).value in ("一致", "不一致")]
print("照合シート:", ws["B4"].value, ws["H4"].value, "| rows", len(res), "不一致", res.count("不一致"))
for r in range(5, ws.max_row + 1):
    if ws.cell(r, 8).value == "不一致": print("  照合 NG:", [ws.cell(r, c).value for c in range(1, 9)])
