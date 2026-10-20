#!/usr/bin/env python3
# カンツギ 計算詳細版の裏取りブックを作る。数字はすべて「前提」を参照する式。tora_v2.js の値は「照合」「崩れたとき」に貼る。
# Usage: python3 build_xlsx.py <workdir with inputs.json & node_values.json> <out.xlsx>
import json
import math
import sys

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule, FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as CL
from openpyxl.worksheet.page import PageMargins
from openpyxl.worksheet.pagebreak import Break
from openpyxl.worksheet.properties import PageSetupProperties

WORK, OUT = sys.argv[1], sys.argv[2]
IN = json.load(open(WORK + "/inputs.json", encoding="utf-8"))
NV = json.load(open(WORK + "/node_values.json", encoding="utf-8"))

FN = "BIZ UDPゴシック"
NAVY, BLUE, RED, GREEN = "1B365D", "2F6DB5", "C62828", "2E7D32"
INK, GRAY, PASTE = "000000", "6B6B6B", "555555"
LIGHT, LINE, PALEG = "E8EDF4", "C9D1DC", "E3F1E4"

F0 = '#,##0;"▲"#,##0;0'
F1 = '#,##0.0;"▲"#,##0.0;0.0'
F2 = '#,##0.00;"▲"#,##0.00;0.00'
F3 = '#,##0.000;"▲"#,##0.000;0.000'
F4 = '#,##0.0000;"▲"#,##0.0000;0.0000'
F6 = '#,##0.000000;"▲"#,##0.000000;0'
P0 = '0%;"▲"0%;0%'
P1 = '0.0%;"▲"0.0%;0.0%'
P2 = '0.00%;"▲"0.00%;0.00%'
YR = '0"年目"'
GEN = "General"

thin = Side(style="thin", color=LINE)
BORDER = Border(bottom=thin)
HDR_FILL = PatternFill("solid", fgColor=NAVY)
SEC_FILL = PatternFill("solid", fgColor=LIGHT)

REG = {}
PENDING = []
t = None


def _ev(me, fx, *args):
    """年ごとのシートの式を、そのシート（me）を t にして作る"""
    global t
    old, t = t, me
    try:
        return fx(*args)
    finally:
        t = old
CF_NEG = {}  # sheet -> list of ranges for red negatives


def font(color=INK, bold=False, size=10, italic=False):
    return Font(name=FN, size=size, bold=bold, italic=italic, color=color)


def reg(key, sheet, row, col):
    assert key not in REG, key
    REG[key] = (sheet, row, col)


def A(key, y=1, here=None):
    sheet, row, col = REG[key]
    s = f"${CL(col + y - 1)}${row}"
    return s if here == sheet else f"'{sheet}'!{s}"


def RNG(key, n=10, here=None, start=1):
    sheet, row, col = REG[key]
    s = f"${CL(col + start - 1)}${row}:${CL(col + start + n - 2)}${row}"
    return s if here == sheet else f"'{sheet}'!{s}"


def text_units(s):
    return sum(1.85 if ord(ch) > 0x2E7F else 1 for ch in str(s))


def fit_height(ws, r, text, width, size=10):
    lines = 0
    for part in str(text).split("\n"):
        lines += max(1, math.ceil(text_units(part) * (size / 10) / max(width - 1, 1)))
    h = (13.5 if size <= 10 else 16) * lines + 4
    cur = ws.row_dimensions[r].height or 15
    ws.row_dimensions[r].height = max(cur, h, 15)


COLOR = {"input": BLUE, "paste": PASTE, "green": GREEN, "calc": INK, "gray": GRAY, "navy": NAVY, "red": RED, "white": "FFFFFF"}


def put(ws, r, c, v, fmt=None, kind="calc", bold=False, align=None, wrap=False, width=None, size=10, fill=None, border=True, indent=0):
    cell = ws.cell(r, c)
    if callable(v):
        PENDING.append((cell, v))
    elif v is not None:
        cell.value = v
    cell.font = font(COLOR[kind], bold, size)
    if fmt:
        cell.number_format = fmt
    cell.alignment = Alignment(horizontal=align, vertical="center", wrap_text=wrap, indent=indent)
    if border:
        cell.border = BORDER
    if fill:
        cell.fill = fill
    if wrap and width and isinstance(v, str):
        fit_height(ws, r, v, width, size)
    return cell


def hdr(ws, r, labels, start=1, height=None):
    for i, t in enumerate(labels):
        if t is False:
            continue
        c = ws.cell(r, start + i)
        if callable(t):
            PENDING.append((c, t))
        elif t is not None:
            c.value = t
        c.font = font("FFFFFF", True)
        c.fill = HDR_FILL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = height or 30


def page(ws, title_rows="3:3"):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr = PageSetupProperties(fitToPage=True)
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    if title_rows:
        ws.print_title_rows = title_rows
    ws.oddHeader.left.text = "カンツギ 計算詳細版 裏取り"
    ws.oddHeader.left.size = 8
    ws.oddHeader.left.font = FN + ",Regular"
    ws.oddFooter.center.font = FN + ",Regular"
    ws.oddHeader.center.text = "&A"
    ws.oddHeader.center.size = 11
    ws.oddHeader.center.font = FN + ",Bold"
    ws.oddFooter.center.text = "&P / &N"
    ws.oddFooter.center.size = 8
    ws.page_margins = PageMargins(left=0.4, right=0.4, top=0.75, bottom=0.6, header=0.3, footer=0.3)
    ws.print_options.horizontalCentered = True


def setup(ws, widths, title, note):
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[CL(i)].width = w
    n = len(widths)
    put(ws, 1, 1, title, kind="navy", bold=True, size=14, border=False)
    ws.row_dimensions[1].height = 24
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=n)
    put(ws, 2, 1, note, kind="gray", wrap=True, width=sum(widths), border=False)
    ws.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="top")


def sec_row(ws, r, title, ncol):
    for c in range(1, ncol + 1):
        ws.cell(r, c).fill = SEC_FILL
    put(ws, r, 1, title, kind="navy", bold=True, size=10.5, fill=SEC_FILL, border=False)
    ws.row_dimensions[r].height = 19


def note_row(ws, r, text, ncol, widths, kind="gray"):
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncol)
    put(ws, r, 1, text, kind=kind, wrap=True, width=sum(widths), border=False)
    ws.cell(r, 1).alignment = Alignment(wrap_text=True, vertical="top")


wb = Workbook()
wb.remove(wb.active)
SHEETS = ["使い方", "前提", "時給", "お客様の効果", "値段と原価", "導入支援", "体制と人件費", "拠点と売上", "損益と現金",
          "黒字ライン", "LTV・CAC", "崩れたとき", "資金と黒字倒産", "市場規模", "照合"]
for n in SHEETS:
    wb.create_sheet(n)


# ======================================================================
# 前提
# ======================================================================
class Pre:
    name = "前提"
    W = [6, 42, 15, 7] + [10] * 10 + [64]

    def __init__(self):
        ws = self.ws = wb[self.name]
        setup(ws, self.W, "前提：入力はこのシートだけ",
              "青字＝入力（ここだけを書き換える）、黒字＝計算。他のシートはすべてこのシートを参照する式。年ごとの項目は「値」の列が1年目で、右へ2〜10年目。"
              "1〜5年目は tora_v2.js の ROLES・GROWTH、6〜10年目は EXT の値。確度：統計＝公的統計／制度＝法律・料率など制度で決まる値／公表＝メーカー・クラウドの公表価格／見積＝社内見積／仮説＝実証で測る／仮定＝計算のために置いた値／換算＝単位の換算。")
        hdr(ws, 3, ["番号", "項目", "単位", "確度", "値（年ごとは1年目）"] + [f"{y}年目" for y in range(2, 11)] + ["出典・メモ"])
        ws.freeze_panes = "E4"
        page(ws)
        self.r = 4

    def sec(self, title):
        sec_row(self.ws, self.r, title, 15)
        self.r += 1

    def sub(self, labels, first="項目"):
        hdr(self.ws, self.r, [None, first, None, None] + labels + [None] * (10 - len(labels)) + [None], height=28)
        self.r += 1

    def row(self, no, key, label, unit, conf, vals, src, fmt=F0):
        ws, r = self.ws, self.r
        put(ws, r, 1, no, align="center", kind="gray")
        put(ws, r, 2, label, wrap=True, width=self.W[1])
        put(ws, r, 3, unit, kind="gray", wrap=True, width=self.W[2])
        put(ws, r, 4, conf, align="center", bold=(conf == "仮説"), kind=("navy" if conf == "仮説" else "calc"))
        put(ws, r, 15, src, kind="gray", wrap=True, width=self.W[14], size=9)
        reg(key, self.name, r, 5)
        if not isinstance(vals, list):
            vals = [vals]
        for i, v in enumerate(vals):
            c = 5 + i
            if callable(v):
                f = v
                put(ws, r, c, (lambda f=f, r=r, c=c: f(r, c)), fmt, "calc")
            elif isinstance(v, str) and v.startswith("="):
                put(ws, r, c, v, fmt, "input")  # 入力を式で書いたもの（例 =1/3）
            else:
                put(ws, r, c, v, fmt, "input")
        self.r += 1


P = Pre()
Ph = "前提"
pa = lambda k, y=1: A(k, y, here=Ph)
prev_col = lambda r, c: f"={CL(c - 1)}{r}"

P.sec("0. 単位の換算（前提ではなく、式の中に数字を書かないために置く）")
P.row("0-1", "c_min", "1時間の分", "分/時間", "換算", 60, "")
P.row("0-2", "c_month", "1年の月数", "か月/年", "換算", 12, "")
P.row("0-3", "c_yen", "1万円の円", "円/万円", "換算", 10000, "")
P.row("0-4", "c_day", "1年の日数", "日/年", "換算", 365, "")
P.row("0-5", "c_oku", "1億円の万円", "万円/億円", "換算", 10000, "")
P.row("0-6", "c_mtok", "AIの単価の基準（100万トークン）", "トークン", "換算", 1000000, "AIの単価は100万トークンあたりのドルで書く")

P.sec("1. 時給（お客様の1時間の人件費）")
P.row("1-1", "wage_e", "職長（電気工事従事者）の年収", "万円/年", "統計", IN["WAGE"]["electrician"],
      "厚生労働省「賃金構造基本統計調査」令和5年 電気工事従事者：きまって支給する給与 月36.74万円×12＋賞与等110.06万円＝550.94万円（残業代を含む）。令和6年も約548〜551万円")
P.row("1-2", "wage_m", "施工管理者の年収", "万円/年", "統計", IN["WAGE"]["manager"],
      "同調査で代用：建築技術者 約642万円（令和6年）と土木技術者 約604万円（令和5年）の間。施工管理技士は標本が少ないので使わない")
P.row("1-3", "emp_load", "会社負担（社会保険料）を含める倍率", "倍", "制度", IN["WAGE"]["employerLoad"],
      "会社負担 約17%：健康保険 約5.0%＋厚生年金9.15%＋雇用保険（建設）1.1%＋子ども・子育て拠出金0.36%＋労災 約1.2%", fmt=F2)
P.row("1-4", "hours_m", "建設業の月の総実労働時間（残業を含む）", "時間/月", "統計", 161.5,
      "厚生労働省「毎月勤労統計調査」令和6年 建設業の総実労働時間（所定内＋所定外）", fmt=F1)
P.row("1-5", "hours_y", "年間総実労働時間", "時間/年", "統計", lambda r, c: f"={CL(c)}{r - 1}*{pa('c_month')}",
      "月161.5時間×12＝1,938時間（tora_v2.js の annualHours）")
P.row("1-6", "ot", "残業の割増", "倍", "制度", IN["WAGE"]["overtimePremium"], "労働基準法：時間外労働は25%以上の割増。参考の計算だけに使い、計画には入れない", fmt=F2)
P.row("1-7", "ref_h", "参考：所定の1日の時間", "時間/日", "仮定", 8, "所定の目安（1日8時間）。参考の比較だけに使う")
P.row("1-8", "ref_d", "参考：所定の年の日数（目安）", "日/年", "仮定", 244, "所定の目安（年約244日）。8時間×244日＝1,952時間は所定の目安で、実績ではない")

P.sec("2. 端末と値段（1拠点）")
P.row("2-1", "devices", "1拠点の台数", "台/拠点", "見積", IN["PLAN"]["devices"], "1班に1台、職長が掛ける。5現場×2班（計算詳細版 p.10）")
P.row("2-2", "workdays", "端末を使う日数", "日/年", "見積", IN["PLAN"]["workDays"], "上限に近い想定。現場の稼働日数で減る")
P.row("2-3", "base_fee", "基盤の利用料", "万円/年/拠点", "見積", IN["PLAN"]["baseFee"], "月5万円×12")
P.row("2-4", "dev_fee", "端末ライセンス", "万円/年/台", "見積", IN["PLAN"]["deviceFee"], "月2万円×12。録画は端末に無制限、AIで解析する場面に月の上限")
P.row("2-5", "setup_fee", "導入支援（初年度のみ）", "万円/拠点", "見積", IN["PLAN"]["setup"], "42人日の原価141万円の約1.4倍（他社の相場は未確認）")
P.row("2-6", "glass_price", "カメラ付きAIグラスの本体", "万円/台", "公表", 8, "Ray-Ban Meta 7.4〜8.9万円")
P.row("2-7", "disp_price", "表示付きARグラスの本体", "万円/台", "公表", 12, "Meta Ray-Ban Display $799")
P.row("2-8", "lease_rate", "リースの月額料率（3年リース）", "割合/月", "見積", 0.03, "3年リースの料率 約3%／月。リース会社の見積で確認する", fmt=P1)
P.row("2-9", "glass_n", "構成：カメラ付きAIグラス", "台/拠点", "見積", IN["PLAN"]["mix"]["glass"], "記録用")
P.row("2-10", "disp_n", "構成：表示付きARグラス", "台/拠点", "見積", IN["PLAN"]["mix"]["display"], "手元に図面を出す班")
P.row("2-11", "lease_glass", "リースの年額：カメラ付き", "万円/年/台", "公表", lambda r, c: f"={pa('glass_price')}*{pa('lease_rate')}*{pa('c_month')}",
      "本体8万円×月額料率3%×12か月", fmt=F2)
P.row("2-12", "lease_disp", "リースの年額：表示付き", "万円/年/台", "公表", lambda r, c: f"={pa('disp_price')}*{pa('lease_rate')}*{pa('c_month')}",
      "本体12万円×月額料率3%×12か月", fmt=F2)
P.row("2-13", "mix_sum", "構成の台数の合計（確認）", "台/拠点", "見積", lambda r, c: f"={pa('glass_n')}+{pa('disp_n')}",
      "2-1の台数と同じであること。リースはお客様の支払いだけに効き、当社の売上・原価には入らない")

P.sec("3. 削減時間（1台1日）と手戻り（参考）")
P.row("3-1", "save_e", "職長の削減（日報・撮影・黒板）", "分/台/日", "仮説", IN["SAVE"]["electricianMin"], "実証で導入前2週間と導入後4週間を比べて測る")
P.row("3-2", "save_m", "施工管理者の削減（写真整理・写真台帳）", "分/台/日", "仮説", IN["SAVE"]["managerMin"], "同上")
P.row("3-3", "rw_sales", "参考：お客様の完成工事高", "万円/年", "仮定", 300000, "中堅100人規模の目安30億円。計画に入れない")
P.row("3-4", "rw_rate", "参考：手戻りが工事費に占める割合", "割合", "仮定", 0.02, "海外研究の範囲0.4〜6%、RC造の事例1.85〜2.1%から控えめに", fmt=P1)
P.row("3-5", "rw_target", "参考：ARが効く対象工種の割合", "割合", "仮説", 0.3, "", fmt=P0)
P.row("3-6", "rw_cut", "参考：手戻りの削減率", "割合", "仮説", 0.2, "", fmt=P0)
P.sub(["1", "2", "3", "4"], first="感度表の軸")
P.row("3-7", "sens_days", "感度表：使う日数", "日/年", "見積", [180, 200, 220, 240], "240日は上限に近い。現場の稼働で減る場合を見る")
P.row("3-8", "sens_e", "感度表：職長の削減（25分・35分・60分の内訳）", "分/台/日", "仮説", [10, 15, 25], "tora_v2.js の感度表と同じ内訳（10＋15、15＋20、25＋35）")
P.row("3-9", "sens_m", "感度表：施工管理者の削減（同上）", "分/台/日", "仮説", [15, 20, 35], "")

P.sec("4. AI・保存の原価（1台・月）とデータ基盤")
P.row("4-1", "fx", "為替", "円/ドル", "仮定", 150, "$1＝150円")
P.row("4-2", "rec_h", "1日の録画時間", "時間/日", "見積", 4, "")
P.row("4-3", "rec_d", "1か月の稼働日", "日/月", "見積", 20, "")
P.row("4-4", "stt_price", "音声の文字起こしの単価", "ドル/分", "公表", 0.003, "クラウドの文字起こしの公表単価の目安（多めに見た上限）", fmt=F3)
P.row("4-5", "frames", "大事な場面の画像", "枚/台/月", "見積", 3000, "録画は端末に全部残し、選んだ場面だけを送る")
P.row("4-6", "model_tok", "モデルの置き方：画像1枚のトークン", "トークン/枚", "仮定", 1000, "tora_v2.js の CLOUD.vision（3,000枚×約1,000トークンで月1,000円）")
P.row("4-7", "img_w", "画像の幅", "ピクセル", "公表", 1280, "開発者向けの映像は720p（計算詳細版 p.26）")
P.row("4-8", "img_h", "画像の高さ", "ピクセル", "公表", 720, "")
P.row("4-9", "img_div", "画像のトークン数の換算（幅×高さ÷この値）", "ピクセル/トークン", "公表", 750, "Claude の画像のトークン数の目安＝幅×高さ÷750")
P.row("4-10", "prompt_tok", "画像1枚あたりの指示文", "トークン/枚", "見積", 300, "")
P.row("4-11", "out_tok", "画像1枚あたりの出力", "トークン/枚", "見積", 150, "")
P.row("4-12", "haiku_in", "Claude Haiku 5.5 の入力単価", "ドル/100万トークン", "公表", 0.10, "Anthropic の API 価格（10万トークン以下の要求）", fmt=F2)
P.row("4-13", "haiku_out", "Claude Haiku 5.5 の出力単価", "ドル/100万トークン", "公表", 0.50, "同上", fmt=F2)
P.row("4-14", "sonnet_in", "Claude Sonnet 5.5 の入力単価", "ドル/100万トークン", "公表", 2, "Anthropic の API 価格", fmt=F2)
P.row("4-15", "sonnet_out", "Claude Sonnet 5.5 の出力単価", "ドル/100万トークン", "公表", 10, "同上", fmt=F2)
P.row("4-16", "rep_n", "報告書の件数", "件/台/日", "見積", 1, "1日1件")
P.row("4-17", "rep_in", "報告書1件の入力", "トークン/件", "見積", 20000, "Sonnet 5.5 で作るとして検算する")
P.row("4-18", "rep_out", "報告書1件の出力", "トークン/件", "見積", 2000, "")
P.row("4-19", "clip_gb", "クラウドへ送る切り出し動画", "GB/台/月", "見積", 10, "1日4時間の録画そのものは端末に残し、送らない")
P.row("4-20", "keep_m", "クラウドの保存期間", "か月", "見積", 36, "3年保存")
P.row("4-21", "s3", "保存の単価（Amazon S3 相当）", "ドル/GB・月", "公表", 0.025, "", fmt=F3)
P.row("4-22", "cloud_stt", "モデルの値：音声の文字起こし", "円/台/月", "見積", IN["CLOUD"]["transcribe"], "tora_v2.js CLOUD.transcribe")
P.row("4-23", "cloud_vis", "モデルの値：画像解析", "円/台/月", "見積", IN["CLOUD"]["vision"], "tora_v2.js CLOUD.vision")
P.row("4-24", "cloud_rep", "モデルの値：報告書の生成", "円/台/月", "見積", IN["CLOUD"]["report"], "tora_v2.js CLOUD.report")
P.row("4-25", "cloud_sto", "モデルの値：動画の保存", "円/台/月", "見積", IN["CLOUD"]["storage"], "tora_v2.js CLOUD.storage（3年間の平均で置いた値）")
P.row("4-26", "cloud_ops", "モデルの値：転送・監視・バックアップ・端末管理（MDM）・余裕", "円/台/月", "見積", IN["CLOUD"]["ops"], "tora_v2.js CLOUD.ops")
P.row("4-27", "full", "全編を送り毎秒解析した場合の原価", "円/台/月", "見積", int(round(IN["FULL"] / 12 * 10000)), "tora_v2.js FULL_UPLOAD_PER_DEVICE（5万円×12）。動画1日4時間・毎秒1枚の解析と保存")
P.row("4-28", "infra_m", "データ基盤（会社ごとの保管庫）", "万円/月/拠点", "見積", 1, "tora_v2.js siteInfra（年12万円）")
P.row("4-29", "infra_y", "データ基盤の年額", "万円/年/拠点", "見積", lambda r, c: f"={CL(c)}{r - 1}*{pa('c_month')}", "月1万円×12")
P.row("4-30", "over_k", "追加料金＝原価の何倍か", "倍", "仮定", 3, "解析量のプラン案（参考）だけに使う")
P.row("4-31", "over_frames", "追加の単位：画像", "枚", "仮定", 1000, "")
P.row("4-32", "over_gb", "追加の単位：保存", "GB/台/月", "仮定", 10, "")
P.row("4-33", "round_unit", "追加料金の丸めの単位", "円", "仮定", 1000, "")
P.sub(["1", "2", "3"], first="粗利率の表の軸")
P.row("4-34", "margin_devs", "粗利率の表：1拠点の台数", "台/拠点", "仮定", [5, 10, 25], "計算詳細版 p.12")

P.sec("5. 導入支援の作業（1拠点）")
for i, (task, days) in enumerate(IN["SETUP_TASKS"], 1):
    P.row(f"5-{i}", f"task{i}", task, "人日/拠点", "見積", days, "tora_v2.js SETUP_TASKS")
P.row("5-7", "setup_travel", "導入支援の交通費", "万円/拠点", "見積", IN["setupTravel"], "16回訪問×約1万円")
P.row("5-8", "cs_days", "導入コンサルタントの年間稼働日", "日/年", "見積", 240, "日額＝年収600万円×1.2÷240日（tora_v2.js CSDAY）")
P.row("5-9", "cs_cap", "導入コンサルタント1人が回せる新規拠点", "拠点/人/年", "見積", 6, "2か月で1拠点（tora_v2.js GROWTH.newSites の注）")

P.sec("6. 体制（年収と期末の人数。1〜5年目＝ROLES、6〜10年目＝EXT.heads）")
P.sub(["1年目"] + [f"{y}年目" for y in range(2, 11)], first="役割")
ROLES = IN["ROLES"]
for i, role in enumerate(ROLES):
    name, sal = role[0], role[1]
    if isinstance(sal, list):
        vals = sal[:5] + [prev_col] * 5
        src = "見積。段階的に上げ、6年目以降は5年目と同じ"
    else:
        vals = [sal] + [prev_col] * 9
        src = "見積（相場は要確認）。2年目以降は前の年を参照"
    P.row(f"6-{i + 1}", f"sal{i}", f"年収：{name}", "万円/年", "見積", vals, src)
for i, role in enumerate(ROLES):
    heads = role[2] + IN["EXT"]["heads"][i]
    P.row(f"6-{i + 9}", f"head{i}", f"人数：{role[0]}", "人（期末）", "見積", heads, role[3])
P.row("6-17", "load", "当社の会社負担を含める倍率", "倍", "見積", IN["LOAD"], "法定福利16%＋福利厚生・研修4%", fmt=F2)
P.row("6-18", "agent_fee", "人材紹介の手数料（年収に対する割合）", "割合", "見積", 0.3, "人材紹介の相場 年収の30%", fmt=P0)
P.row("6-19", "agent_share", "紹介会社を使う割合", "割合", "見積", 0.5, "", fmt=P0)
P.row("6-20", "recruit_rate", "採用費率", "割合", "見積", lambda r, c: f"={pa('agent_fee')}*{pa('agent_share')}", "年収30%×紹介会社経由50%＝15%（tora_v2.js RECRUIT_RATE）", fmt=P1)
P.row("6-21", "adv_tech", "社外の技術顧問（大学・映像AI）", "万円/月", "見積", 10, "")
P.row("6-22", "adv_ind", "社外の業界顧問（電気工事業のOB）", "万円/月", "見積", 5, "")
P.row("6-23", "advisors", "社外顧問の年額", "万円/年", "見積", lambda r, c: f"=({pa('adv_tech')}+{pa('adv_ind')})*{pa('c_month')}", "(月10万＋月5万)×12＝180万円（tora_v2.js ADVISORS）")

P.sec("7. 成長（拠点・個社設定・経費・資金。1〜5年目＝GROWTH、6〜10年目＝EXT）")
P.sub(["1年目"] + [f"{y}年目" for y in range(2, 11)])
G, EXT = IN["GROWTH"], IN["EXT"]
P.row("7-1", "new_sites", "新規拠点", "拠点/年", "見積", G["newSites"] + EXT["newSites"],
      "導入コンサルタント1人＝年6拠点の上限から。6年目以降は関東の中堅で頭打ち（電線工場・海外は入れない）")
P.row("7-2", "churn", "解約率（期首の拠点に対して）", "割合/年", "見積", G["churnRate"], "B2B SaaS の年間解約は大企業5〜10%・中小30〜50%。乗り換えにくさを見て中間より低め", fmt=P0)
P.row("7-3", "new_share", "新規の初年度の課金割合（半年分）", "割合", "仮定", 0.5, "新規は初年度に半年分を払う", fmt=P0)
P.row("7-4", "trial_year", "試験導入の年（利用料なし）", "年目", "仮定", 1, "1年目の3拠点は有償の試験導入。利用料は入らず、原価は3か月分かかる", fmt=YR)
P.row("7-5", "trial_m", "試験導入で原価がかかる月数", "か月", "仮定", 3, "")
P.row("7-6", "trial_share", "試験導入の原価の割合", "割合", "仮定", lambda r, c: f"={CL(c)}{r - 1}/{pa('c_month')}", "3か月÷12か月", fmt=P0)
P.row("7-7", "contract", "個社設定の売上", "万円/年", "見積", G["contract"] + EXT["contract"], "その会社の帳票・既存システムとの連携の設定。件数×単価の仮置き（売上の5%以下）")
P.row("7-8", "contract_margin", "個社設定の粗利率", "割合", "見積", "=1/3", "人件費（2/3）を原価に振り替える", fmt=P1)
OTHER = [("office", "その他の経費：オフィス（賃料＋共益費）", "川崎 15坪→40坪、坪1.25〜1.3万円"),
         ("ads", "その他の経費：広告・展示会", "展示会・業界団体・Web・営業資料"),
         ("travel", "その他の経費：営業・定着支援の交通費", "導入支援の交通費は原価"),
         ("tools", "その他の経費：開発用の端末・検証機・ソフト", ""),
         ("experts", "その他の経費：専門家（税理士・弁護士・弁理士・社労士）", ""),
         ("misc", "その他の経費：雑費", ""),
         ("ml", "その他の経費：AIの学習・評価の計算費", "クラウドGPU 月15〜30万円、2年目から")]
for j, (k, label, src) in enumerate(OTHER):
    P.row(f"7-{9 + j}", f"other_{k}", label, "万円/年", "見積", G["other"][k] + EXT["other"][k], src)
P.row("7-16", "funding", "資金投入（自己資金）", "万円/年", "見積", G["funding"] + EXT["funding"],
      "課題の前提「豊富な資金」。1年目2億円、3年目に撤退ラインを越えていれば5,000万円")
P.row("7-17", "prep_fund", "準備期の資金（試作品・実証）", "万円", "見積", 500, "準備期に使い切る。1年目期首の現金には入れない（tora_v2.js と同じ）")
P.row("7-18", "open_sites", "開業時の拠点", "拠点", "仮定", 0, "")
P.row("7-19", "open_cash", "1年目期首の現金", "万円", "仮定", 0, "tora_v2.js は現金0から始め、資金投入と営業損益を足す")
P.row("7-20", "min_m", "最低残高の目安（販管費の何か月分）", "か月", "見積", 3, "")

P.sec("8. シナリオのつまみ（この初期値のとき tora_v2.js の計画と一致する）")
P.row("8-1", "k_price", "値段の倍率（基盤と端末ライセンス）", "倍", "仮定", 1, "tora_v2.js run({priceK})", fmt=F2)
P.row("8-2", "k_churn", "解約率の上書き（空なら7-2を使う）", "割合/年", "仮定", [None], "tora_v2.js run({churnRate})", fmt=P0)
P.row("8-3", "k_cost", "原価の倍率（AI・保存の1台あたり）", "倍", "仮定", 1, "tora_v2.js run({perDevice})。データ基盤には掛けない", fmt=F2)
P.row("8-4", "k_salary", "年収の上乗せ（全員）", "万円/人", "仮定", 0, "tora_v2.js run({salaryAdd})。採用費には効かない")
P.row("8-5", "k_hire", "採用連動のk（空なら連動しない）", "倍", "仮定", [None], "tora_v2.js run({hireK})：社員数＝4＋(計画−4)×k を四捨五入", fmt=F2)
P.row("8-6", "hire_base", "採用連動のときに残す人数", "人", "仮定", 4, "tora_v2.js の式の4")

P.sec("9. LTV・CAC と黒字ライン")
P.row("9-1", "ltv_plan", "計画で使う継続年数", "年", "仮定", 7, "解約が年10%なら平均10年。7年は控えめ")
P.row("9-2", "cac_year", "CACを計算する年", "年目", "仮定", 5, "tora_v2.js unit()：5年目の営業3人・広告・新規", fmt=YR)
P.row("9-3", "bl_year", "台数の黒字ラインを計算する年", "年目", "仮定", 5, "", fmt=YR)
P.row("9-4", "check_devs", "黒字ラインで確かめる台数", "台/拠点", "仮定", 15, "計算詳細版 p.18・21")
P.sub(["1", "2", "3", "4"], first="表の軸")
P.row("9-5", "ltv_years", "LTVの表：継続年数", "年", "仮定", [3, 5, 7, 10], "")
P.row("9-6", "cac_k", "CACの表：倍率", "倍", "仮定", [1, 2, 4], "獲得が難しくなった場合")

P.sec("10. 資金と黒字倒産")
P.row("10-1", "ar_days", "売掛の回収（月額の後払い・支払いサイト）", "日", "仮定", 60, "")
P.row("10-2", "hire_lead", "採用の先行（拠点が増える前に入社）", "か月", "仮定", 3, "")
P.row("10-3", "dev_start", "端末込みの料金を始める年", "年目", "仮定", 3, "計算詳細版 p.29（3年目から端末込みも選べる）", fmt=YR)
P.row("10-4", "dev_unit", "端末1台の調達額", "万円/台", "公表", lambda r, c: f"={pa('glass_price')}", "カメラ付きAIグラス（2-6）。表示付きを含む構成なら1拠点 8×8＋2×12＝88万円")
P.row("10-5", "dev_recover", "端末代を料金で回収する年数（4年まで）", "年", "仮定", 3, "リースと同じ3年で、買った翌年から均等に回収")
P.row("10-6", "surge", "急な拠点増（新規が計画より何割多いか）", "割合", "仮定", 0.5, "", fmt=P0)
P.row("10-7", "deposit", "導入支援の着手金の割合", "割合", "見積", 0.5, "計算詳細版 p.13（着手金50%）", fmt=P0)

P.sec("11. 市場規模（tora_market.js の SEG と SOM）")
P.row("11-1", "mk_total", "電気工事業の許可業者", "社", "統計", 65497, "国土交通省「建設業許可業者数調査」令和7年3月末")
P.sub(["顧客数（社）", "1社の年額（万円/年）", "ターゲット（1＝入れる）"], first="業種・規模")
SEG = IN["SEG"]
for i, (ind, size, n, price, insam, conf, note) in enumerate(SEG):
    if ind == "電気工事会社" and size.startswith("小"):
        nval = lambda r, c: f"={pa('mk_total')}-{pa('seg3')}-{pa('seg4')}"
        cf = "統計−仮定"
        note = "65,497社−大300−中2,100（残りを小規模とする）。" + note
    else:
        nval, cf = n, "仮定"
    P.row(f"11-{i + 2}", f"seg{i}", f"{ind}・{size}", "社・万円/年・1/0", cf, [nval, price, 1 if insam else 0], note)
P.sub(["地域の割合", "会える率", "成約率", "1社の拠点数"], first="SOMの漏斗")
SOM = IN["SOM"]
for j, (ci, region, reach, close, per, note) in enumerate(SOM):
    P.row(f"11-{10 + j}", f"som{j}", f"{SEG[ci][0]}・{SEG[ci][1]}", "割合・拠点/社", "仮説", [region, reach, close, per], note, fmt=F2)
P.row("11-12", "light_devs", "参考：小規模向けライト版の台数", "台/拠点", "仮定", 3, "基盤＋3台。計画に入れない")
P.row("11-13", "light_share", "参考：小規模のうちライト版で取れる割合", "割合", "仮定", [0.001, 0.005, 0.01], "上振れの目安。計画に入れない", fmt=P1)

P.sec("12. 照合")
P.row("12-1", "tol", "一致とみなす差（絶対値）", "各行の単位", "仮定", 0.000001, "Excel と JavaScript はどちらも倍精度。差は計算順による丸め誤差だけ", fmt=F6)


# ======================================================================
# 共通：年ごとのシート（TS）と、値を縦に並べるシート（SS）
# ======================================================================
class TS:
    W = [38, 11] + [10.5] * 10 + [60]

    def __init__(self, name, title, note):
        ws = self.ws = wb[name]
        self.name = name
        setup(ws, self.W, title, note)
        hdr(ws, 3, ["項目", "単位"] + [None] * 10 + ["式（言葉）"])
        for y in range(1, 11):
            c = ws.cell(3, 2 + y, y)
            c.number_format = YR
        ws.freeze_panes = "C4"
        page(ws)
        self.r = 4
        self.first = 4

    def Y(self, y):
        return f"{CL(2 + y)}$3"

    def a(self, key, y):
        return A(key, y, here=self.name)

    def rng(self, key, n=10, start=1):
        return RNG(key, n, here=self.name, start=start)

    def sec(self, title):
        sec_row(self.ws, self.r, title, 13)
        self.r += 1

    def row(self, key, label, unit, f, desc, fmt=F0, bold=False, f1=None, kind="calc", years=range(1, 11)):
        ws, r = self.ws, self.r
        put(ws, r, 1, label, bold=bold, wrap=True, width=self.W[0])
        put(ws, r, 2, unit, kind="gray", wrap=True, width=self.W[1])
        put(ws, r, 13, desc, kind="gray", wrap=True, width=self.W[12] - 2, size=9, indent=1)
        if key:
            reg(key, self.name, r, 3)
        for y in range(1, 11):
            if y not in years:
                put(ws, r, 2 + y, "─", align="right", kind="gray")
                continue
            fx = f1 if (y == 1 and f1 is not None) else f
            put(ws, r, 2 + y, (lambda fx=fx, y=y, me=self: "=" + _ev(me, fx, y)), fmt, kind, bold)
        self.r += 1
        return r

    def res(self, key, label, unit, f, desc, fmt=F0, kind="calc", bold=True):
        ws, r = self.ws, self.r
        if kind == "green":
            kind = "calc"
            ws.conditional_formatting.add(f"C{r}", FormulaRule(formula=[f"ISNUMBER(C{r})"], font=Font(color=GREEN, bold=True)))
            ws.conditional_formatting.add(f"C{r}", FormulaRule(formula=[f"NOT(ISNUMBER(C{r}))"], font=Font(color=RED, bold=True)))
        put(ws, r, 1, label, bold=bold, wrap=True, width=self.W[0])
        put(ws, r, 2, unit, kind="gray")
        put(ws, r, 3, (lambda me=self: "=" + _ev(me, f)), fmt, kind, bold, align="right")
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=13)
        put(ws, r, 4, desc, kind="gray", wrap=True, width=sum(self.W[3:13]), size=9, indent=1)
        if key:
            reg(key, self.name, r, 3)
        self.r += 1
        return r

    def note(self, text, kind="gray"):
        note_row(self.ws, self.r, text, 13, self.W, kind)
        self.r += 1

    def neg_red(self):
        CF_NEG.setdefault(self.name, []).append(f"C4:L{self.r}")

    def pbreak(self):
        self.ws.row_breaks.append(Break(id=self.r - 1))


class SS:
    W = [40, 13, 14, 14, 14, 46]

    def __init__(self, name, title, note, W=None):
        if W:
            self.W = W
        ws = self.ws = wb[name]
        self.name = name
        setup(ws, self.W, title, note)
        hdr(ws, 3, ["項目", "単位", "値", None, None, "式（言葉）"])
        ws.freeze_panes = "A4"
        page(ws)
        self.r = 4

    def a(self, key, y=1):
        return A(key, y, here=self.name)

    def sec(self, title):
        sec_row(self.ws, self.r, title, 6)
        self.r += 1

    def row(self, key, label, f, unit, desc, fmt=F0, bold=False, kind="calc"):
        ws, r = self.ws, self.r
        put(ws, r, 1, label, bold=bold, wrap=True, width=self.W[0])
        put(ws, r, 2, unit, kind="gray", wrap=True, width=self.W[1])
        put(ws, r, 3, (lambda: "=" + f()), fmt, kind, bold, align="right")
        ws.merge_cells(start_row=r, start_column=4, end_row=r, end_column=6)
        put(ws, r, 4, desc, kind="gray", wrap=True, width=sum(self.W[3:6]), size=9, indent=1)
        if key:
            reg(key, self.name, r, 3)
        self.r += 1
        return r

    def mhdr(self, labels, height=34):
        hdr(self.ws, self.r, labels, height=height)
        self.r += 1

    def mrow(self, cells, desc=None, desc_from=None, key=None):
        """cells: list of (value_or_callable, fmt, kind, bold) from column A"""
        ws, r = self.ws, self.r
        for i, cell in enumerate(cells):
            vals = list(cell) if isinstance(cell, tuple) else [cell]
            v, fmt, kind, bold = vals + [None, None, "calc", False][len(vals):]
            if callable(v):
                vv = (lambda v=v: "=" + v())
            else:
                vv = v
            put(ws, r, 1 + i, vv, fmt, kind or "calc", bold, align=("right" if (i and kind != "gray") else "left"), wrap=(i in (0, 1)), width=self.W[i])
        if desc is not None:
            c0 = desc_from or (len(cells) + 1)
            if c0 < 6:
                ws.merge_cells(start_row=r, start_column=c0, end_row=r, end_column=6)
            put(ws, r, c0, desc, kind="gray", wrap=True, width=sum(self.W[c0 - 1:6]), size=9, indent=1)
        if key:
            reg(key, self.name, r, 3)
        self.r += 1
        return r

    def note(self, text, kind="gray"):
        note_row(self.ws, self.r, text, 6, self.W, kind)
        self.r += 1

    def neg_red(self):
        CF_NEG.setdefault(self.name, []).append(f"C4:E{self.r}")

    def pbreak(self):
        self.ws.row_breaks.append(Break(id=self.r - 1))


# ======================================================================
# 時給
# ======================================================================
ws = wb["時給"]
WJ = [40, 15, 15, 12, 66]
setup(ws, WJ, "時給：お客様の1時間の人件費", "1時間の人件費＝年収×会社負担÷年間総実労働時間。年収も総実労働時間も残業を含むので、分子と分母がそろっている（計算詳細版 p.5）。")
hdr(ws, 3, ["項目", "職長（電工）", "施工管理者", "単位", "式（言葉）"])
ws.freeze_panes = "A4"
page(ws)
HJ = "時給"
jr = [4]


def jrow(key, label, fe, fm, unit, desc, fmt=F0, bold=False, kind="calc"):
    r = jr[0]
    put(ws, r, 1, label, bold=bold, wrap=True, width=WJ[0])
    put(ws, r, 2, (lambda: "=" + fe()) if fe else None, fmt, kind, bold, align="right")
    put(ws, r, 3, (lambda: "=" + fm()) if fm else None, fmt, kind, bold, align="right")
    put(ws, r, 4, unit, kind="gray")
    put(ws, r, 5, desc, kind="gray", wrap=True, width=WJ[4], size=9)
    if key:
        reg(key, HJ, r, 2)
    jr[0] += 1


ja = lambda k, y=1: A(k, y, here=HJ)
jrow("j_wage", "年収（残業代・賞与を含む）", lambda: A("wage_e"), lambda: A("wage_m"), "万円/年", "賃金構造基本統計調査（前提1-1・1-2）")
jrow("j_load", "会社負担を含める倍率", lambda: A("emp_load"), lambda: A("emp_load"), "倍", "社会保険料の会社負担 約17%（前提1-3）", fmt=F2)
jrow("j_cost", "会社負担込みの年の人件費", lambda: f"{ja('j_wage')}*{ja('j_load')}", lambda: f"{ja('j_wage', 2)}*{ja('j_load', 2)}", "万円/年", "年収×会社負担", fmt=F2)
jrow("j_hours", "年間総実労働時間（残業を含む）", lambda: A("hours_y"), lambda: A("hours_y"), "時間/年", "毎月勤労統計 建設業 月161.5時間×12（前提1-5）")
jrow("j_h_man", "1時間の人件費", lambda: f"{ja('j_cost')}/{ja('j_hours')}", lambda: f"{ja('j_cost', 2)}/{ja('j_hours', 2)}", "万円/時間", "年収×会社負担÷年間総実労働時間", fmt=F6)
jrow("j_h", "1時間の人件費（円）", lambda: f"{ja('j_h_man')}*{A('c_yen')}", lambda: f"{ja('j_h_man', 2)}*{A('c_yen')}", "円/時間", "上の万円×10,000", fmt=F2, bold=True)
jrow("j_m", "1分の人件費", lambda: f"{ja('j_h')}/{A('c_min')}", lambda: f"{ja('j_h', 2)}/{A('c_min')}", "円/分", "1時間の人件費÷60", fmt=F2)
r = jr[0]
sec_row(ws, r, "参考（計画には入れない）", 5)
jr[0] += 1
jrow("j_ot", "削った時間がすべて残業だった場合の1時間の価値", lambda: f"{ja('j_h')}*{A('ot')}", lambda: f"{ja('j_h', 2)}*{A('ot')}", "円/時間", "1時間の人件費×残業の割増1.25", fmt=F0)
jrow("j_ot_eff", "そのときのお客様の効果（1拠点）", lambda: f"{A('docs')}*{A('ot')}", None, "万円/年", "お客様の効果（年約499万円）×1.25。計画には入れない", fmt=F2)
jrow("j_ref_hours", "参考：所定の目安の時間（1日8時間×244日）", lambda: f"{A('ref_h')}*{A('ref_d')}", lambda: f"{A('ref_h')}*{A('ref_d')}", "時間/年", "1日8時間×約244日。所定の目安で、実績ではない")
jrow("j_ref_rate", "参考：所定の目安で割った1時間", lambda: f"{ja('j_cost')}/{ja('j_ref_hours')}*{A('c_yen')}", lambda: f"{ja('j_cost', 2)}/{ja('j_ref_hours', 2)}*{A('c_yen')}", "円/時間", "年収×会社負担÷1,952時間。値は近いが意味が違うので使わない", fmt=F0)
jr[0] += 1
for t in [
    "注1：総実労働時間は毎月勤労統計の「総実労働時間」で、所定外（残業）を含む実績。",
    "注2：年収（賃金構造基本統計調査）も、きまって支給する給与（残業代を含む）と賞与の合計。分子にも分母にも残業が入っている。",
    "注3：残業を除いた所定内の時間で割ると、分母だけが小さくなり単価が高く出るので使わない。「1日8時間×約244日＝1,952時間」は所定の目安で実績ではない（値は1,938時間に近いが意味が違う）。",
    "注4：削った時間がすべて残業なら、割増1.25倍の分だけ効果は大きくなる（年約624万円）。計画には入れない。お客様の効果は控えめに見ている。当社の売上と利益は、この単価には左右されない。",
]:
    note_row(ws, jr[0], t, 5, WJ)
    jr[0] += 1
CF_NEG.setdefault(HJ, []).append(f"B4:C{jr[0]}")


# ======================================================================
# お客様の効果
# ======================================================================
cv = SS("お客様の効果", "お客様の効果：1拠点10台で年約499万円、支払いの約1.5倍",
        "効果はAIグラスを掛けた班の分だけで数える（社員全員が減るとは数えない）。お客様の支払いは当社への利用料とグラスのリース。tora_v2.js の customerValue() と同じ式（計算詳細版 p.6〜8）。")
c = cv.a
cv.sec("1台1日の効果")
cv.row("cv_e", "職長の削減", lambda: A("save_e"), "分/台/日", "日報・撮影・黒板（前提3-1・仮説）")
cv.row("cv_m", "施工管理者の削減", lambda: A("save_m"), "分/台/日", "写真整理・写真台帳（前提3-2・仮説）")
cv.row("cv_hw", "職長の1時間の人件費", lambda: A("j_h"), "円/時間", "時給シート", fmt=F2)
cv.row("cv_hm", "施工管理者の1時間の人件費", lambda: A("j_h", 2), "円/時間", "時給シート", fmt=F2)
cv.row("per_day_yen", "1台1日の効果", lambda: f"{c('cv_e')}/{A('c_min')}*{c('cv_hw')}+{c('cv_m')}/{A('c_min')}*{c('cv_hm')}", "円/台/日",
       "15分÷60×職長の時給＋20分÷60×施工管理者の時給", fmt=F2, bold=True)
cv.row("per_day_man", "1台1日の効果（万円）", lambda: f"{c('per_day_yen')}/{A('c_yen')}", "万円/台/日", "上の円÷10,000", fmt=F6)
cv.sec("1拠点・1年の効果")
cv.row("cv_dev", "1拠点の台数", lambda: A("devices"), "台/拠点", "前提2-1")
cv.row("cv_days", "端末を使う日数", lambda: A("workdays"), "日/年", "前提2-2。上限に近い想定")
cv.row("docs", "お客様の効果（年）", lambda: f"{c('per_day_man')}*{c('cv_dev')}*{c('cv_days')}", "万円/年", "1台1日の効果×台数×日数", fmt=F2, bold=True)
cv.row("cv_min", "1台1日の削減の合計", lambda: f"{c('cv_e')}+{c('cv_m')}", "分/台/日", "職長15分＋施工管理者20分")
cv.row("freed_min", "浮く時間（分）", lambda: f"{c('cv_min')}*{c('cv_dev')}*{c('cv_days')}", "分/年", "35分×10台×240日")
cv.row("freed_h", "浮く時間（時間）", lambda: f"{c('freed_min')}/{A('c_min')}", "時間/年", "上の分÷60", fmt=F1)
cv.sec("お客様の支払い（1拠点・2年目以降の年額）")
cv.row("pay_base", "当社：基盤", lambda: A("base_k"), "万円/年", "月5万円×12（値段と原価。値段の倍率を反映）")
cv.row("pay_dev", "当社：端末ライセンス×台数", lambda: f"{A('dev_k')}*{c('cv_dev')}", "万円/年", "月2万円×12×10台")
cv.row("pay_lic", "当社への利用料", lambda: f"{c('pay_base')}+{c('pay_dev')}", "万円/年", "基盤＋端末ライセンス")
cv.row("pay_lg", "リース：カメラ付きAIグラス", lambda: f"{A('glass_n')}*{A('lease_glass')}", "万円/年", "8台×(8万円×3%×12)", fmt=F2)
cv.row("pay_ld", "リース：表示付きARグラス", lambda: f"{A('disp_n')}*{A('lease_disp')}", "万円/年", "2台×(12万円×3%×12)", fmt=F2)
cv.row("pay_lease", "リースの合計", lambda: f"{c('pay_lg')}+{c('pay_ld')}", "万円/年", "お客様からリース会社へ払う（当社の売上ではない）", fmt=F2)
cv.row("cost", "お客様の支払いの合計", lambda: f"{c('pay_lic')}+{c('pay_lease')}", "万円/年", "利用料＋リース", fmt=F2, bold=True)
cv.row("ratio", "倍率（効果÷支払い）", lambda: f"{c('docs')}/{c('cost')}", "倍", "お客様の効果÷支払いの合計", fmt=F4, bold=True)
cv.pbreak()
cv.sec("元が取れる分（1台1日に何分減れば支払いと同じになるか）")
cv.row("pay_per", "1台1日の支払い", lambda: f"{c('cost')}/({c('cv_dev')}*{c('cv_days')})*{A('c_yen')}", "円/台/日", "支払い÷(台数×日数)", fmt=F2)
cv.row("per_min", "同じ配分での1分の人件費", lambda: f"({c('cv_hw')}*{c('cv_e')}/{c('cv_min')}+{c('cv_hm')}*{c('cv_m')}/{c('cv_min')})/{A('c_min')}", "円/分",
       "(職長の時給×15/35＋施工管理者の時給×20/35)÷60", fmt=F4)
cv.row("be", "元が取れる分", lambda: f"{c('pay_per')}/{c('per_min')}", "分/台/日", "1台1日の支払い÷同じ配分での1分の人件費", fmt=F2, bold=True)

cv.sec("感度表：使う日数×1台1日の削減")
sens_lab = lambda i: (lambda: f'TEXT({A("sens_e", i)}+{A("sens_m", i)},"0")&"分"&CHAR(10)&"（"&{A("sens_e", i)}&"＋"&{A("sens_m", i)}&"）"')
cv.mhdr(["年の効果\n使う日数＼1台1日の削減（職長＋施工管理）", "単位", (lambda: "=" + sens_lab(1)()), (lambda: "=" + sens_lab(2)()), (lambda: "=" + sens_lab(3)()), "式（言葉）"])
for d in range(1, 5):
    cells = [((lambda d=d: f"{A('sens_days', d)}"), '0"日"', "calc", True), ("万円/年", None, "gray", False)]
    for i in range(1, 4):
        cells.append(((lambda d=d, i=i: f"({A('sens_e', i)}/{A('c_min')}*{c('cv_hw')}+{A('sens_m', i)}/{A('c_min')}*{c('cv_hm')})/{A('c_yen')}*{c('cv_dev')}*{A('sens_days', d)}"), F1, "calc", False))
    cv.mrow(cells, "(分÷60×時給)の合計×台数×日数÷10,000", desc_from=6, key=f"sens_eff{d}")
cv.mhdr(["倍率\n使う日数＼1台1日の削減（職長＋施工管理）", "単位", (lambda: "=" + sens_lab(1)()), (lambda: "=" + sens_lab(2)()), (lambda: "=" + sens_lab(3)()), "式（言葉）"])
for d in range(1, 5):
    cells = [((lambda d=d: f"{A('sens_days', d)}"), '0"日"', "calc", True), ("倍", None, "gray", False)]
    for i in range(1, 4):
        cells.append(((lambda d=d, i=i: f"{A('sens_eff' + str(d), i)}/{c('cost')}"), F2, "calc", False))
    cv.mrow(cells, "年の効果÷お客様の支払い（331.68万円）", desc_from=6, key=f"sens_ratio{d}")
cv.mhdr(["元が取れる分\n使う日数", "単位", "元が取れる分", "式（言葉）", False, False])
cv.ws.merge_cells(start_row=cv.r - 1, start_column=4, end_row=cv.r - 1, end_column=6)
for d in range(1, 5):
    cells = [((lambda d=d: f"{A('sens_days', d)}"), '0"日"', "calc", True), ("分/台/日", None, "gray", False),
             ((lambda d=d: f"{c('cost')}/({c('cv_dev')}*{A('sens_days', d)})*{A('c_yen')}/{c('per_min')}"), F1, "calc", False)]
    cv.mrow(cells, "支払い÷(台数×日数)÷同じ配分での1分の人件費（削減の量には左右されない）", desc_from=4, key=f"sens_be{d}")
cv.note("注：240日は上限に近い想定。雨天・現場の切れ目・端末を使わない日があれば、現場の稼働日数に合わせて減る。日数が減ると効果は比例して減り、元が取れる分は増える。"
        "25分と60分の内訳は tora_v2.js の感度表と同じ（10＋15、25＋35）。")

cv.pbreak()
cv.sec("手戻りの削減（参考：計画にも提案の金額にも入れない）")
cv.row("rw", "手戻りの削減（参考）", lambda: f"{A('rw_sales')}*{A('rw_rate')}*{A('rw_target')}*{A('rw_cut')}", "万円/年", "完成工事高30億円×手戻り2%×対象工種30%×削減20%", bold=True)
cv.row("ratio_up", "手戻りを足した場合の倍率（参考）", lambda: f"({c('docs')}+{c('rw')})/{c('cost')}", "倍", "(お客様の効果＋手戻り)÷支払い", fmt=F2)
cv.note("注：手戻り率は各社の現場事情で大きく違い、検証の費用も見合わない可能性があるため、計画にも提案の金額にも入れない。")
cv.note("注：「月6.5時間（書類電子化の実績）」は原典を確認できておらず、ARグラスの効果とも違うので使わない。")
cv.neg_red()


# ======================================================================
# 値段と原価
# ======================================================================
pc = SS("値段と原価", "値段と原価：1拠点300万円、原価84万円、粗利216万円（72%）",
        "1台の原価は、クラウド・AIの公表単価から積み上げて検算した。計画の数字は tora_v2.js の CLOUD の値（モデルの値）を使い、検算の値と並べる（計算詳細版 p.9〜12）。")
p = pc.a
pc.sec("1拠点の売上（年）")
pc.row("base_k", "基盤（値段の倍率を反映）", lambda: f"{A('base_fee')}*{A('k_price')}", "万円/年/拠点", "月5万円×12×値段の倍率")
pc.row("dev_k", "端末ライセンス1台（値段の倍率を反映）", lambda: f"{A('dev_fee')}*{A('k_price')}", "万円/年/台", "月2万円×12×値段の倍率")
pc.row("pc_dev", "1拠点の台数", lambda: A("devices"), "台/拠点", "前提2-1")
pc.row("price_site", "1拠点の利用料", lambda: f"{p('base_k')}+{p('dev_k')}*{p('pc_dev')}", "万円/年/拠点", "基盤60＋端末24×10台", bold=True)
pc.row("price_site0", "参考：値段の倍率を掛ける前の1拠点の利用料", lambda: f"{A('base_fee')}+{A('dev_fee')}*{p('pc_dev')}", "万円/年/拠点", "市場規模の計算に使う")

pc.sec("1拠点の原価と粗利（年）")
pc.row("perdev0", "1台の原価（年）", lambda: f"{A('total_model', 1)}*{A('c_month')}/{A('c_yen')}", "万円/年/台", "6,000円×12か月÷10,000（6,000円は下の内訳の合計）", fmt=F2)
pc.row("perdev", "1台の原価（原価の倍率を反映）", lambda: f"{p('perdev0')}*{A('k_cost')}", "万円/年/台", "1台の原価×原価の倍率（前提8-3）", fmt=F2)
pc.row("pc_infra", "データ基盤", lambda: A("infra_y"), "万円/年/拠点", "月1万円×12")
pc.row("cogs_site", "1拠点の原価", lambda: f"{p('pc_infra')}+{p('perdev')}*{p('pc_dev')}", "万円/年/拠点", "データ基盤12＋1台7.2×10台", fmt=F2, bold=True)
pc.row("gp_site", "1拠点の粗利", lambda: f"{p('price_site')}-{p('cogs_site')}", "万円/年/拠点", "1拠点の利用料−1拠点の原価", fmt=F2, bold=True)
pc.row("margin_site", "粗利率（利用料）", lambda: f"{p('gp_site')}/{p('price_site')}", "割合", "1拠点の粗利÷1拠点の利用料", fmt=P1, bold=True)

pc.sec("粗利率と台数・使い方（1拠点）")
pc.mhdr(["台数", "単位", "1拠点の利用料\n（万円/年）", (lambda: f'="粗利率"&CHAR(10)&"大事な場面だけ"&CHAR(10)&"月"&TEXT({A("total_model", 1)},"#,##0")&"円/台"'),
         (lambda: f'="粗利率"&CHAR(10)&"全編を毎秒解析"&CHAR(10)&"月"&TEXT({A("full")},"#,##0")&"円/台"'), "式（言葉）"], height=48)
for i in range(1, 4):
    pc.mrow([((lambda i=i: A("margin_devs", i)), '0"台"', "calc", True), ("万円/年・割合", None, "gray", False),
             ((lambda i=i: f"{p('base_k')}+{p('dev_k')}*{A('margin_devs', i)}"), F0, "calc", False),
             ((lambda i=i: f"1-({p('pc_infra')}+{p('perdev')}*{A('margin_devs', i)})/({p('base_k')}+{p('dev_k')}*{A('margin_devs', i)})"), P1, "calc", False),
             ((lambda i=i: f"1-({p('pc_infra')}+{A('full')}*{A('c_month')}/{A('c_yen')}*{A('margin_devs', i)})/({p('base_k')}+{p('dev_k')}*{A('margin_devs', i)})"), P1, "calc", False)],
            "1−(データ基盤＋1台の原価×台数)÷(基盤＋端末×台数)", desc_from=6, key=f"mg{i}")
pc.note("注：1台の値段（月2万円）が1台の原価（月0.6万円）を賄うので、台数が増えても粗利率はほぼ同じ。全編を送ると台数が増えるほど赤字が膨らむ。")

pc.pbreak()
pc.sec("1台の原価の内訳（円/台/月）")
pc.mhdr(["項目", "単位", "検算の値", "モデルの値\n（tora_v2.js）", "差\nモデル−検算", "式（言葉）"])
def crow(k, label, f, unit, desc, fmt=F0, model=None, bold=False):
    cells = [(label, None, "calc", bold),
             (unit, None, "gray", False),
             ((lambda f=f: f()) if f else None, fmt, "calc", bold),
             ((lambda model=model: A(model)) if model else None, F0, "calc", bold),
             ((lambda k=k: f"{A(k, 2, here='値段と原価')}-{A(k, 1, here='値段と原価')}") if (model and f) else None, F0, "calc", False)]
    return pc.mrow(cells, desc, desc_from=6, key=k)



crow("stt_min", "音声：録画の分数", lambda: f"{A('rec_h')}*{A('rec_d')}*{A('c_min')}", "分/台/月", "1日4時間×月20日×60分＝80時間")
crow("stt", "音声の文字起こし", lambda: f"{p('stt_min')}*{A('stt_price')}*{A('fx')}", "円/台/月", "4,800分×$0.003×150円。削減時間（35分）とは関係なく、録画した時間に掛かる", F0, "cloud_stt", True)
crow("img_tok", "画像：1枚の画像のトークン", lambda: f"{A('img_w')}*{A('img_h')}/{A('img_div')}", "トークン/枚", "1280×720÷750", F1)
crow("img_in", "画像：1枚の入力トークン", lambda: f"{p('img_tok')}+{A('prompt_tok')}", "トークン/枚", "画像＋指示文300", F1)
crow("hk_frame", "画像解析 Claude Haiku 5.5：1枚", lambda: f"({p('img_in')}*{A('haiku_in')}+{A('out_tok')}*{A('haiku_out')})/{A('c_mtok')}*{A('fx')}", "円/枚",
     "(入力×$0.10＋出力150×$0.50)÷100万×150円", F4)
crow("hk_month", "画像解析 Claude Haiku 5.5：月", lambda: f"{p('hk_frame')}*{A('frames')}", "円/台/月", "1枚×3,000枚", F1)
crow("sn_frame", "画像解析 Claude Sonnet 5.5：1枚", lambda: f"({p('img_in')}*{A('sonnet_in')}+{A('out_tok')}*{A('sonnet_out')})/{A('c_mtok')}*{A('fx')}", "円/枚",
     "(入力×$2＋出力150×$10)÷100万×150円", F4)
crow("sn_month", "画像解析 Claude Sonnet 5.5：月", lambda: f"{p('sn_frame')}*{A('frames')}", "円/台/月", "1枚×3,000枚", F1)
crow("md_frame", "画像解析 モデルの置き方：1枚", lambda: f"{A('cloud_vis')}/{A('frames')}", "円/枚", "月1,000円÷3,000枚（1枚約1,000トークンとして置いた値）", F4)
crow("vis", "画像解析（大事な場面 月3,000枚）", None, "円/台/月", "モデルの1,000円は Haiku 5.5（約100円）と Sonnet 5.5（約2,050円）の間", F0, "cloud_vis", True)
crow("vis_up", "全部を Sonnet 5.5 にした場合の増え方", lambda: f"{p('sn_month')}-{A('cloud_vis')}", "円/台/月", "Sonnet 5.5 の月額−モデルの1,000円", F0)
crow("rep_cost", "報告書：1件の費用（Sonnet 5.5）", lambda: f"({A('rep_in')}*{A('sonnet_in')}+{A('rep_out')}*{A('sonnet_out')})/{A('c_mtok')}*{A('fx')}", "円/件",
     "(入力2万×$2＋出力2千×$10)÷100万×150円", F2)
crow("rep", "報告書の生成", lambda: f"{p('rep_cost')}*{A('rep_n')}*{A('rec_d')}", "円/台/月", "1件×1日1件×20日", F0, "cloud_rep", True)
crow("sto_gb", "保存：3年間の平均の保管量", lambda: f"{A('clip_gb')}*{A('keep_m')}/2", "GB/台", "月10GB×36か月÷2（直線で増える量の平均）", F0)
crow("sto", "動画の保存", lambda: f"{p('sto_gb')}*{A('s3')}*{A('fx')}", "円/台/月", "平均180GB×$0.025×150円", F0, "cloud_sto", True)
crow("ops", "転送・監視・バックアップ・端末管理（MDM）・余裕", None, "円/台/月", "見積（内訳の検算はしていない）", F0, "cloud_ops", True)
r_tot = pc.mrow([("合計（計画で使う値）", None, "calc", True), ("円/台/月", None, "gray", False), (None, None), ((lambda: f"{A('stt', 2)}+{A('vis', 2)}+{A('rep', 2)}+{A('sto', 2)}+{A('ops', 2)}"), F0, "calc", True), (None, None)],
                "文字起こし＋画像解析＋報告書＋保存＋転送など（モデルの値の合計）", desc_from=6)
REG["total_model"] = ("値段と原価", r_tot, 4)
crow("tot_h", "参考：検算の合計（画像は Haiku 5.5）", lambda: f"{p('stt')}+{p('hk_month')}+{p('rep')}+{p('sto')}+{A('cloud_ops')}", "円/台/月", "検算の値の合計（転送などはモデルの値）", F0)
crow("tot_s", "参考：検算の合計（画像は Sonnet 5.5）", lambda: f"{p('stt')}+{p('sn_month')}+{p('rep')}+{p('sto')}+{A('cloud_ops')}", "円/台/月", "同上", F0)
pc.pbreak()
pc.sec("広い原価で見た粗利率（5年目）")
y5 = 5
pc.row("b_sales", "5年目の売上", lambda: A("sales", y5), "万円/年", "拠点と売上")
pc.row("b_cogs", "5年目の原価（今の分類）", lambda: A("cogs", y5), "万円/年", "利用料の原価＋導入支援の原価＋個社設定の人件費", fmt=F1)
pc.row("b_gp", "5年目の粗利（今の分類）", lambda: A("gp", y5), "万円/年", "売上−原価", fmt=F1)
pc.row("b_rate", "粗利率（今の分類）", lambda: f"{p('b_gp')}/{p('b_sales')}", "割合", "粗利÷売上", fmt=P1, bold=True)
pc.row("b_cs", "導入コンサルタントの人件費", lambda: A("pay4", y5), "万円/年", "8人×600万円×1.2（体制と人件費。計画の人数・年収）", fmt=F1)
pc.row("b_tr", "うち導入支援と個社設定へ振り替えた分", lambda: A("transfer", y5), "万円/年", "損益と現金の振替", fmt=F1)
pc.row("b_rest", "振り替えていない分（定着支援・サポート）", lambda: f"{p('b_cs')}-{p('b_tr')}", "万円/年", "導入後の定着支援と問い合わせ対応の時間", fmt=F1)
pc.row("b_sv", "現場監修の人件費（ラベル付け・品質の確認）", lambda: A("pay3", y5), "万円/年", "2人×650万円×1.2", fmt=F1)
pc.row("b_move", "原価へ移す額", lambda: f"{p('b_rest')}+{p('b_sv')}", "万円/年", "振り替えていない分＋現場監修", fmt=F1)
pc.row("b_gp2", "広い原価での粗利", lambda: f"{p('b_gp')}-{p('b_move')}", "万円/年", "粗利−原価へ移す額", fmt=F1)
pc.row("b_rate2", "広い原価での粗利率", lambda: f"{p('b_gp2')}/{p('b_sales')}", "割合", "SaaS の慣行どおり、定着支援・サポート・現場監修を原価に入れた場合", fmt=P1, bold=True)
pc.row("b_lic", "5年目の利用料", lambda: A("license", y5), "万円/年", "拠点と売上")
pc.row("b_lic_cost", "5年目の利用料の原価", lambda: A("lic_cost", y5), "万円/年", "損益と現金", fmt=F1)
pc.row("b_lic_rate", "利用料の粗利率（今の分類）", lambda: f"1-{p('b_lic_cost')}/{p('b_lic')}", "割合", "1−利用料の原価÷利用料", fmt=P1)
pc.row("b_lic_rate2", "利用料の粗利率（広い原価）", lambda: f"({p('b_lic')}-{p('b_lic_cost')}-{p('b_move')})/{p('b_lic')}", "割合", "移す額をすべて利用料の原価にした場合：(利用料−利用料の原価−原価へ移す額)÷利用料", fmt=P1)
pc.row("b_sga2", "販管費（広い分類）", lambda: f"{A('sga', y5)}-{p('b_move')}", "万円/年", "販管費−原価へ移す額", fmt=F1)
pc.row("b_op2", "営業利益（広い分類）", lambda: f"{p('b_gp2')}-{p('b_sga2')}", "万円/年", "広い粗利−広い販管費", fmt=F1)
pc.row("b_opdiff", "今の分類の営業利益との差", lambda: f"{p('b_op2')}-{A('op', y5)}", "万円/年", "0になる：分類を変えても営業利益は変わらない", fmt=F1, bold=True)
pc.note("注：分類を変えても営業利益は変わらない。サポート（問い合わせ窓口）は導入コンサルタントが受け持つので「振り替えていない分」に含めた。人件費は計画の人数・年収（つまみを掛ける前）。")

pc.pbreak()
pc.sec("保存：4年目以降の定常（3年保存を続けた場合）")
pc.mhdr(["項目", "単位", "検算の値", "モデルの値\n（tora_v2.js）", "差\nモデル−検算", "式（言葉）"])
crow("sto_ss_gb", "保存：4年目以降の定常の保管量", lambda: f"{A('clip_gb')}*{A('keep_m')}", "GB/台", "月10GB×36か月", F0)
crow("sto_ss", "保存：定常の費用", lambda: f"{p('sto_ss_gb')}*{A('s3')}*{A('fx')}", "円/台/月", "360GB×$0.025×150円", F0)
crow("sto_diff", "保存：定常とモデルの差（1台）", lambda: f"{p('sto_ss')}-{A('cloud_sto')}", "円/台/月", "定常−モデルの700円", F0)
crow("sto_diff_site", "保存：定常での1拠点の増え方", lambda: f"{p('sto_diff')}*{p('pc_dev')}*{A('c_month')}/{A('c_yen')}", "万円/年/拠点", "差×10台×12か月÷10,000", F1)
crow("sto_diff_pt", "保存：定常での粗利率の下がり幅", lambda: f"{p('sto_diff_site')}/{p('price_site')}", "ポイント", "1拠点の増え方÷1拠点の利用料", P1)
pc.note("注：モデルは保存を3年間の平均で置いている。定常（4年目以降）では1台 月+650円（1拠点 年+7.8万円、粗利率−2.6ポイント）になる。")

pc.sec("解析量のプラン案（参考・計画に入れない）")
pc.row("std_frames", "標準に含む画像", lambda: A("frames"), "枚/台/月", "前提4-5")
pc.row("std_gb", "標準に含む保存", lambda: A("clip_gb"), "GB/台/月", "前提4-19")
pc.row("of_cost_s", "追加1,000枚の原価（Sonnet 5.5）", lambda: f"{p('sn_frame')}*{A('over_frames')}", "円/月", "1枚の Sonnet 5.5 の費用×1,000枚", fmt=F1)
pc.row("of_cost_h", "追加1,000枚の原価（Haiku 5.5・参考）", lambda: f"{p('hk_frame')}*{A('over_frames')}", "円/月", "1枚の Haiku 5.5 の費用×1,000枚", fmt=F1)
pc.row("of_price", "追加1,000枚ごとの料金", lambda: f"ROUND({p('of_cost_s')}*{A('over_k')}/{A('round_unit')},0)*{A('round_unit')}", "円/台/月", "Sonnet 5.5 の原価×3を千円単位に丸める", bold=True)
pc.row("of_margin", "その粗利率", lambda: f"1-{p('of_cost_s')}/{p('of_price')}", "割合", "1−原価÷料金", fmt=P1)
pc.row("og_cost", "追加10GB/月の原価（3年保存の平均）", lambda: f"{A('over_gb')}*{A('keep_m')}/2*{A('s3')}*{A('fx')}", "円/月", "毎月10GB多く送り3年保存：平均180GB×$0.025×150円", fmt=F0)
pc.row("og_cost_ss", "追加10GB/月の原価（4年目以降の定常）", lambda: f"{A('over_gb')}*{A('keep_m')}*{A('s3')}*{A('fx')}", "円/月", "360GB×$0.025×150円", fmt=F0)
pc.row("og_price", "追加10GBごとの料金", lambda: f"ROUND({p('og_cost')}*{A('over_k')}/{A('round_unit')},0)*{A('round_unit')}", "円/台/月", "3年平均の原価×3を千円単位に丸める", bold=True)
pc.row("og_margin", "その粗利率（3年平均）", lambda: f"1-{p('og_cost')}/{p('og_price')}", "割合", "1−原価÷料金", fmt=P1)
pc.row("og_margin_ss", "その粗利率（定常）", lambda: f"1-{p('og_cost_ss')}/{p('og_price')}", "割合", "定常でも料金が原価を上回る", fmt=P1)
pc.note("注：標準は1台 月3,000枚・保存10GBまで。超えた分は1,000枚ごと・10GBごとに追加料金。追加料金は原価の約3倍で置き、計算してから千円単位に丸めた。")
pc.neg_red()


# ======================================================================
# 導入支援
# ======================================================================
su = SS("導入支援", "導入支援：42人日＋交通費で原価141万円、値段200万円、粗利59万円",
        "導入支援の人件費は原価に振り替える（販管費から引く）。tora_v2.js の SETUP_TASKS・setupDays・CSDAY・setupTravel と同じ（計算詳細版 p.13）。")
s = su.a
su.sec("作業と人日（1拠点）")
for i, (task, days) in enumerate(IN["SETUP_TASKS"], 1):
    su.row(f"su_t{i}", task, (lambda i=i: A(f"task{i}")), "人日", f"前提5-{i}")
su.row("setup_days", "合計（約2か月）", lambda: f"SUM({s('su_t1')}:{s('su_t6')})", "人日", "6つの作業の合計", bold=True)
su.sec("原価と粗利（1拠点）")
su.row("su_sal", "導入コンサルタントの年収", lambda: A("sal4", 1), "万円/年", "前提6-5（1年目の値。tora_v2.js は600を直接使う）")
su.row("su_load", "会社負担を含める倍率", lambda: A("load"), "倍", "前提6-17", fmt=F2)
su.row("su_days", "年間稼働日", lambda: A("cs_days"), "日/年", "前提5-8")
su.row("csday", "1人日の人件費", lambda: f"{s('su_sal')}*{s('su_load')}/{s('su_days')}", "万円/人日", "年収600万円×1.2÷240日", fmt=F2)
su.row("setup_labor", "人件費", lambda: f"{s('setup_days')}*{s('csday')}", "万円/拠点", "42人日×3万円", fmt=F1)
su.row("su_travel", "交通費", lambda: A("setup_travel"), "万円/拠点", "16回訪問×約1万円")
su.row("setup_cost", "原価", lambda: f"{s('setup_labor')}+{s('su_travel')}", "万円/拠点", "人件費＋交通費", fmt=F1, bold=True)
su.row("su_price", "値段", lambda: A("setup_fee"), "万円/拠点", "前提2-5")
su.row("setup_gp", "粗利", lambda: f"{s('su_price')}-{s('setup_cost')}", "万円/拠点", "値段−原価", fmt=F1, bold=True)
su.row("setup_rate", "粗利率", lambda: f"{s('setup_gp')}/{s('su_price')}", "割合", "粗利÷値段", fmt=P1)
su.sec("お客様の支払い（1拠点）")
su.row("cust_y1", "お客様の1年目の支払い", lambda: f"{A('pay_lic')}+{A('pay_lease')}+{s('su_price')}", "万円", "利用料300＋リース31.68＋導入支援200", fmt=F2, bold=True)
su.row("cust_y2", "お客様の2年目以降の支払い", lambda: A("cost"), "万円/年", "利用料300＋リース31.68", fmt=F2)
su.row("our_rev", "そのうち当社の売上（利用料）", lambda: A("pay_lic"), "万円/年", "1年目は導入支援200万円も当社の売上")
su.note("注：当社の売上は利用料300万円（初年度は導入支援200万円も）。リース31.7万円はお客様からリース会社へ払うもので、当社の売上ではない。")
su.note("注：導入コンサルタント1人が回せる新規拠点は年6（2か月で1拠点）。新規拠点がこの上限内かは「体制と人件費」で確かめる。")
su.neg_red()
su.ws.page_setup.fitToHeight = 1  # 1ページに収める


# ======================================================================
# 体制と人件費
# ======================================================================
tm = TS("体制と人件費", "体制と人件費：8つの役割×年収×人数。5年目は23人",
        "人件費＝年収×人数×1.2（会社負担）。採用費＝増えた人数×年収×15%。tora_v2.js の team() と run() の人件費・採用費と同じ（計算詳細版 p.17）。")
t = tm
NR = len(ROLES)
t.sec("年収（万円/年）")
for i, role in enumerate(ROLES):
    t.row(f"sal{i}_y", f"年収：{role[0]}", "万円/年", (lambda y, i=i: A(f"sal{i}", y)), f"前提6-{i + 1}")
t.sec("人数（期末）")
for i, role in enumerate(ROLES):
    t.row(f"h{i}", f"人数：{role[0]}", "人", (lambda y, i=i: A(f"head{i}", y)), f"前提6-{i + 9}")
t.row("staff_plan", "計画の社員数", "人", lambda y: f"SUM({CL(2 + y)}{REG['h0'][1]}:{CL(2 + y)}{REG['h7'][1]})", "8つの役割の合計", bold=True)
t.sec("人件費（年収×人数×会社負担）")
for i, role in enumerate(ROLES):
    t.row(f"pay{i}", f"人件費：{role[0]}", "万円/年", (lambda y, i=i: f"{t.a(f'sal{i}_y', y)}*{t.a(f'h{i}', y)}*{A('load')}"), "年収×人数×1.2", fmt=F0)
t.row("pay_total", "人件費の合計（計画）", "万円/年", lambda y: f"SUM({CL(2 + y)}{REG['pay0'][1]}:{CL(2 + y)}{REG['pay7'][1]})", "8つの役割の合計", bold=True)
t.sec("採用費（増えた人の年収×採用費率）")
for i, role in enumerate(ROLES):
    t.row(f"hire{i}", f"増員：{role[0]}", "人", (lambda y, i=i: f"MAX(0,{t.a(f'h{i}', y)}-{t.a(f'h{i}', y - 1)})"), "その年の人数−前の年の人数（減ったときは0）。1年目は開業時0人から",
          f1=(lambda y, i=i: f"{t.a(f'h{i}', y)}"))
t.row("hire_amt", "増員×年収の合計", "万円/年", lambda y: f"SUMPRODUCT({CL(2 + y)}{REG['hire0'][1]}:{CL(2 + y)}{REG['hire7'][1]},{CL(2 + y)}{REG['sal0_y'][1]}:{CL(2 + y)}{REG['sal7_y'][1]})",
      "役割ごとの増員×その年の年収を足す")
t.row("recruit_plan", "採用費（計画）", "万円/年", lambda y: f"{t.a('hire_amt', y)}*{A('recruit_rate')}", "増員×年収の合計×15%", fmt=F1)
t.sec("つまみの反映（初期値なら計画のまま）")
t.row("sal_add", "年収の上乗せ分", "万円/年", lambda y: f"{A('k_salary')}*{A('load')}*{t.a('staff_plan', y)}", "上乗せ×1.2×計画の社員数（前提8-4）")
t.row("staff", "社員数（採用連動の後）", "人", lambda y: f"IF(ISBLANK({A('k_hire')}),{t.a('staff_plan', y)},{A('hire_base')}+ROUND(({t.a('staff_plan', y)}-{A('hire_base')})*{A('k_hire')},0))",
      "kが空なら計画のまま。kがあれば 4＋(計画−4)×k を四捨五入", bold=True)
t.row("scale", "人数の倍率", "倍", lambda y: f"{t.a('staff', y)}/{t.a('staff_plan', y)}", "社員数÷計画の社員数", fmt=F3)
t.row("payroll", "人件費（損益で使う）", "万円/年", lambda y: f"({t.a('pay_total', y)}+{t.a('sal_add', y)})*{t.a('scale', y)}", "(人件費の合計＋上乗せ分)×人数の倍率", bold=True)
t.row("recruit", "採用費（損益で使う）", "万円/年", lambda y: f"{t.a('recruit_plan', y)}*{t.a('scale', y)}", "採用費×人数の倍率", fmt=F1, bold=True)
t.sec("導入コンサルタントの確認（新規拠点は 人数×6 が上限）")
t.row("cs_capy", "1年に回せる新規拠点の上限", "拠点", lambda y: f"{t.a('h4', y)}*{A('cs_cap')}", "導入コンサルタントの人数×6")
t.row("cs_new", "新規拠点", "拠点", lambda y: A("new_sites", y), "前提7-1")
t.row("cs_ok", "判定", "─", lambda y: f'IF({t.a("cs_new", y)}<={t.a("cs_capy", y)},"上限内","超える")', "新規拠点≦上限なら「上限内」", fmt=GEN)
t.row("cs_avail", "稼働できる人日", "人日", lambda y: f"{t.a('h4', y)}*{A('cs_days')}", "人数×240日")
t.row("cs_setup", "導入支援に使う人日", "人日", lambda y: f"{t.a('cs_new', y)}*{A('setup_days')}", "新規拠点×42人日")
t.row("cs_contract", "個社設定に使う人日", "人日", lambda y: f"{A('contract_labor', y)}/{A('csday')}", "個社設定の人件費÷1人日の人件費（3万円）", fmt=F1)
t.row("cs_util", "稼働率（導入支援＋個社設定）", "割合", lambda y: f"({t.a('cs_setup', y)}+{t.a('cs_contract', y)})/{t.a('cs_avail', y)}", "使う人日÷稼働できる人日", fmt=P0)
t.note("注：新規拠点はどの年も上限（人数×6）の内に入っている。ただし個社設定もすべて導入コンサルタントが行うとすると、1〜3年目は稼働率が100%を超える。"
       "その分はエンジニアや代表が受け持つ前提になる（計算詳細版 p.15「導入コンサルタントとエンジニアが行う」）。人件費は役割ごとに計上済みなので、損益の数字は変わらない。")
t.neg_red()


# ======================================================================
# 拠点と売上
# ======================================================================
st = TS("拠点と売上", "拠点と売上：期首−解約＋新規＝期末。売上の柱は既存の利用料",
        "1年目は試験導入で利用料なし。新規は初年度に半年分。解約は期首×10%を四捨五入（1年目は0）。tora_v2.js の run({years:10}) と同じ式（計算詳細版 p.14・15）。")
t = st
t.sec("拠点の積み上げ")
t.row("open", "期首の拠点", "拠点", lambda y: t.a("close", y - 1), "前の年の期末。1年目は開業時の拠点（前提7-18）", f1=lambda y: A("open_sites"))
t.row("churn_used", "使う解約率", "割合/年", lambda y: f"IF(ISBLANK({A('k_churn')}),{A('churn')},{A('k_churn')})", "前提8-2が空なら前提7-2", fmt=P0)
t.row("churn_n", "解約", "拠点", lambda y: f"IF({t.Y(y)}=1,0,ROUND({t.a('open', y)}*{t.a('churn_used', y)},0))", "1年目は0。2年目から 期首×解約率 を四捨五入")
t.row("new", "新規", "拠点", lambda y: A("new_sites", y), "前提7-1")
t.row("close", "期末の拠点", "拠点", lambda y: f"{t.a('open', y)}-{t.a('churn_n', y)}+{t.a('new', y)}", "期首−解約＋新規", bold=True)
t.sec("利用料")
t.row("billed_exist", "通年で課金する拠点（期首−解約）", "拠点", lambda y: f"{t.a('open', y)}-{t.a('churn_n', y)}", "期首−解約")
t.row("price_y", "1拠点の利用料", "万円/年", lambda y: A("price_site"), "値段と原価（値段の倍率を反映）")
t.row("new_share_y", "新規の初年度の課金割合", "割合", lambda y: f"IF({t.Y(y)}={A('trial_year')},0,{A('new_share')})", "試験導入の年（1年目）は0、それ以外は半年分", fmt=P0)
t.row("lic_exist", "既存の利用料", "万円/年", lambda y: f"{t.a('billed_exist', y)}*{t.a('price_y', y)}", "(期首−解約)×1拠点の利用料")
t.row("lic_new", "新規の利用料", "万円/年", lambda y: f"{t.a('new', y)}*{t.a('price_y', y)}*{t.a('new_share_y', y)}", "新規×1拠点の利用料×初年度の課金割合")
t.row("license", "利用料の合計", "万円/年", lambda y: f"{t.a('lic_exist', y)}+{t.a('lic_new', y)}", "既存＋新規", bold=True)
t.row("lost", "参考：解約で減った利用料", "万円/年", lambda y: f"{t.a('churn_n', y)}*{t.a('price_y', y)}", "解約×1拠点の利用料")
t.sec("サービスと売上")
t.row("setup_rev", "導入支援", "万円/年", lambda y: f"{t.a('new', y)}*{A('setup_fee')}", "新規×200万円")
t.row("contract_rev", "個社設定", "万円/年", lambda y: A("contract", y), "前提7-7")
t.row("service", "サービスの売上", "万円/年", lambda y: f"{t.a('setup_rev', y)}+{t.a('contract_rev', y)}", "導入支援＋個社設定")
t.row("sales", "売上", "万円/年", lambda y: f"{t.a('license', y)}+{t.a('service', y)}", "利用料＋サービス", bold=True)
t.row("share_exist", "既存の利用料の割合", "割合", lambda y: f"{t.a('lic_exist', y)}/{t.a('sales', y)}", "既存の利用料÷売上", fmt=P0)
t.neg_red()


# ======================================================================
# 損益と現金
# ======================================================================
pl = TS("損益と現金", "損益と現金：10台のままでも6年目に黒字、8年目に累計を取り戻す",
        "販管費＝人件費−原価への振替＋採用費＋その他の経費＋顧問。導入支援と個社設定の人件費は原価へ振り替える。tora_v2.js の run({years:10}) と同じ式（計算詳細版 p.3・18・30）。")
t = pl
t.sec("原価")
t.row("billed", "課金する拠点（通年換算）", "拠点", lambda y: f"{A('billed_exist', y)}+{A('new', y)}*{A('new_share_y', y)}", "(期首−解約)＋新規×初年度の課金割合", fmt=F1)
t.row("cogs_site_y", "1拠点の原価", "万円/年", lambda y: A("cogs_site"), "値段と原価（データ基盤＋1台の原価×台数）", fmt=F1)
t.row("trial_cost", "試験導入の原価（3か月分）", "万円/年", lambda y: f"IF({t.Y(y)}={A('trial_year')},{A('new', y)}*{t.a('cogs_site_y', y)}*{A('trial_share')},0)", "試験導入の年だけ 新規×1拠点の原価×3/12")
t.row("lic_cost", "利用料の原価", "万円/年", lambda y: f"{t.a('billed', y)}*{t.a('cogs_site_y', y)}+{t.a('trial_cost', y)}", "課金する拠点×1拠点の原価＋試験導入の原価")
t.row("setup_labor_y", "導入支援の人件費", "万円/年", lambda y: f"{A('new', y)}*{A('setup_labor')}", "新規×126万円（42人日×3万円）")
t.row("setup_travel_y", "導入支援の交通費", "万円/年", lambda y: f"{A('new', y)}*{A('setup_travel')}", "新規×15万円")
t.row("setup_cost_y", "導入支援の原価", "万円/年", lambda y: f"{t.a('setup_labor_y', y)}+{t.a('setup_travel_y', y)}", "人件費＋交通費")
t.row("contract_labor", "個社設定の人件費", "万円/年", lambda y: f"{A('contract_rev', y)}*(1-{A('contract_margin')})", "個社設定×2/3（粗利1/3）", fmt=F1)
t.row("cogs", "原価の合計", "万円/年", lambda y: f"{t.a('lic_cost', y)}+{t.a('setup_cost_y', y)}+{t.a('contract_labor', y)}", "利用料の原価＋導入支援の原価＋個社設定の人件費", fmt=F1, bold=True)
t.sec("粗利")
t.row("sales_pl", "売上", "万円/年", lambda y: A("sales", y), "拠点と売上")
t.row("gp", "粗利", "万円/年", lambda y: f"{t.a('sales_pl', y)}-{t.a('cogs', y)}", "売上−原価", bold=True)
t.row("gp_rate", "粗利率", "割合", lambda y: f"{t.a('gp', y)}/{t.a('sales_pl', y)}", "粗利÷売上", fmt=P1)
t.row("lic_margin", "利用料の粗利率", "割合", lambda y: f"IF({A('license', y)}=0,0,1-{t.a('lic_cost', y)}/{A('license', y)})", "1−利用料の原価÷利用料（1年目は利用料がないので0）", fmt=P1)
t.sec("販管費")
t.row("payroll_pl", "人件費", "万円/年", lambda y: A("payroll", y), "体制と人件費")
t.row("transfer", "原価への振替", "万円/年", lambda y: f"{t.a('setup_labor_y', y)}+{t.a('contract_labor', y)}", "導入支援の人件費＋個社設定の人件費（販管費から引く）", fmt=F1)
t.row("recruit_pl", "採用費", "万円/年", lambda y: A("recruit", y), "体制と人件費", fmt=F1)
t.row("other", "その他の経費（7項目）", "万円/年", lambda y: "+".join(A(f"other_{k}", y) for k, _, _ in OTHER), "オフィス＋広告＋交通費＋端末・ソフト＋専門家＋雑費＋AIの学習費（前提7-9〜7-15）")
t.row("adv", "社外顧問", "万円/年", lambda y: A("advisors"), "月15万円×12")
t.row("sga", "販管費", "万円/年", lambda y: f"{t.a('payroll_pl', y)}-{t.a('transfer', y)}+{t.a('recruit_pl', y)}+{t.a('other', y)}+{t.a('adv', y)}", "人件費−振替＋採用費＋その他＋顧問", fmt=F1, bold=True)
t.sec("営業利益と現金")
t.row("op", "営業利益", "万円/年", lambda y: f"{t.a('gp', y)}-{t.a('sga', y)}", "粗利−販管費", fmt=F1, bold=True)
t.row("cum", "累計の営業損益", "万円", lambda y: f"{t.a('cum', y - 1)}+{t.a('op', y)}", "前の年の累計＋その年の営業利益", fmt=F1, f1=lambda y: t.a("op", y))
t.row("funding_y", "資金投入", "万円/年", lambda y: A("funding", y), "前提7-16（準備期の500万円は別）")
t.row("cash", "期末の現金", "万円", lambda y: f"{t.a('cash', y - 1)}+{t.a('funding_y', y)}+{t.a('op', y)}", "前の年の現金＋資金投入＋営業利益。1年目は期首0から",
      fmt=F1, bold=True, f1=lambda y: f"{A('open_cash')}+{t.a('funding_y', y)}+{t.a('op', y)}")
t.row("black_flag", "（営業利益が黒字の年）", "年目", lambda y: f'IF({t.a("op", y)}>0,{t.Y(y)},"")', "黒字なら年の番号、赤字なら空", fmt=YR, kind="gray")
t.row("cum_flag", "（累計が黒字の年）", "年目", lambda y: f'IF({t.a("cum", y)}>0,{t.Y(y)},"")', "累計が黒字なら年の番号", fmt=YR, kind="gray")
r_op = REG["op"][1]
t.sec("結果")
t.res("first_black10", "黒字化の年（10年で）", "年目", lambda: f'IF(COUNT({t.rng("black_flag")})=0,"10年でなし",MIN({t.rng("black_flag")}))', "営業利益が初めて黒字になる年", fmt=YR, kind="green")
t.res("first_black5", "黒字化の年（5年で）", "年目", lambda: f'IF(COUNT({t.rng("black_flag", 5)})=0,"5年でなし",MIN({t.rng("black_flag", 5)}))', "1〜5年目だけで見た場合", fmt=YR, kind="green")
t.res("cum_back10", "累計回収の年（10年で）", "年目", lambda: f'IF(COUNT({t.rng("cum_flag")})=0,"10年でなし",MIN({t.rng("cum_flag")}))', "累計の営業損益が初めて黒字になる年", fmt=YR, kind="green")
t.res("min_cash10", "最低の現金（10年）", "万円", lambda: f"MIN({t.rng('cash')})", "期末の現金のいちばん低い値", fmt=F1)
t.res("min_cash_y", "最低の現金の年", "年目", lambda: f"INDEX($C$3:$L$3,1,MATCH({t.a('min_cash10', 1)},{t.rng('cash')},0))", "", fmt=YR)
t.res("min_cash5", "最低の現金（5年）", "万円", lambda: f"MIN({t.rng('cash', 5)})", "1〜5年目", fmt=F1)
t.res("cum_bottom", "累計赤字の底", "万円", lambda: f"MIN({t.rng('cum')})", "累計の営業損益のいちばん低い値", fmt=F1)
t.note("注：簡易計算で、運転資本（売掛・前受）・税・償却を含まない。黒字化の年は緑、マイナスは赤（▲）で示す。")
t.neg_red()
pl.ws.conditional_formatting.add(f"C{r_op}:L{r_op}", FormulaRule(formula=[f"AND(C{r_op}>0,COUNTIF($C{r_op}:C{r_op},\">0\")=1)"],
                                                                  font=Font(color=GREEN, bold=True), fill=PatternFill("solid", fgColor=PALEG, bgColor=PALEG)))


# ======================================================================
# 黒字ライン
# ======================================================================
bl = TS("黒字ライン", "黒字ライン：その年の体制で何拠点・何台あれば黒字か",
        "損益分岐点売上＝販管費÷粗利率。必要な拠点＝(販管費−サービスの粗利)÷1拠点の利用料の粗利216万円（拠点は通年で課金するとして）（計算詳細版 p.19・22）。")
t = bl
t.sec("売上で見る（損益分岐点売上＝販管費÷粗利率）")
t.row("bl_sga", "販管費", "万円/年", lambda y: A("sga", y), "損益と現金", fmt=F1)
t.row("bl_rate", "粗利率", "割合", lambda y: A("gp_rate", y), "損益と現金", fmt=P1)
t.row("bep", "損益分岐点売上", "万円/年", lambda y: f"{t.a('bl_sga', y)}/{t.a('bl_rate', y)}", "販管費÷粗利率", bold=True)
t.row("bl_sales", "実際の売上", "万円/年", lambda y: A("sales", y), "拠点と売上")
t.row("bl_ratio", "分岐点に対する売上の割合", "割合", lambda y: f"{t.a('bl_sales', y)}/{t.a('bep', y)}", "売上÷損益分岐点売上（100%以上で黒字）", fmt=P0)
t.row("safety", "安全余裕率", "割合", lambda y: f'IF({t.a("bl_sales", y)}>{t.a("bep", y)},1-{t.a("bep", y)}/{t.a("bl_sales", y)},"─")', "(売上−分岐点売上)÷売上。赤字の年は「─」", fmt=P0)
t.sec("人で見る")
t.row("bl_staff", "社員数", "人", lambda y: A("staff", y), "体制と人件費（採用連動の後）")
t.row("bl_sites", "期末の拠点", "拠点", lambda y: A("close", y), "拠点と売上")
t.row("sites_per", "社員1人あたりの拠点", "拠点/人", lambda y: f"{t.a('bl_sites', y)}/{t.a('bl_staff', y)}", "期末の拠点÷社員数", fmt=F2, bold=True)
t.row("sga_per", "社員1人あたりの販管費", "万円/人", lambda y: f"{t.a('bl_sga', y)}/{t.a('bl_staff', y)}", "販管費÷社員数")
t.row("need_per", "社員1人あたりに要る課金拠点", "拠点/人", lambda y: f"{t.a('sga_per', y)}/{A('gp_site')}", "社員1人の販管費÷1拠点の利用料の粗利（新規は初年度半額なので期末ではもっと要る）", fmt=F2)
t.sec("その年の体制で黒字に必要な拠点（通年で課金する拠点として）")
t.row("bl_svgp", "サービスの粗利", "万円/年", lambda y: f"{A('service', y)}-{A('setup_cost_y', y)}-{A('contract_labor', y)}", "導入支援＋個社設定−導入支援の原価−個社設定の人件費", fmt=F1)
t.row("bl_sitegp", "1拠点の利用料の粗利", "万円/年", lambda y: A("gp_site"), "値段と原価（300−84）", fmt=F1)
t.row("need_sites", "黒字に必要な課金拠点", "拠点", lambda y: f"({t.a('bl_sga', y)}-{t.a('bl_svgp', y)})/{t.a('bl_sitegp', y)}", "(販管費−サービスの粗利)÷216", fmt=F1, bold=True)
t.row("bl_billed", "実際の課金拠点（通年換算）", "拠点", lambda y: A("billed", y), "損益と現金（1年目は試験導入なので0）", fmt=F1)
t.row("bl_room", "課金拠点の余裕（実際−必要）", "拠点", lambda y: f"{t.a('bl_billed', y)}-{t.a('need_sites', y)}", "マイナス（赤）なら足りない", fmt=F1)
t.sec("5年目に黒字にするのに必要な1拠点の台数（81拠点のまま台数だけ増やす）")
ix = lambda k: f"INDEX({RNG(k, here='黒字ライン')},1,{A('bl_year')})"
t.res("dv_billed", "課金する拠点（通年換算）", "拠点", lambda: ix("bl_billed"), "その年の (期首−解約)＋新規×半年", fmt=F1, bold=False)
t.res("dv_sites", "期末の拠点", "拠点", lambda: ix("bl_sites"), "", bold=False)
t.res("dv_need", "必要な利用料の粗利", "万円/年", lambda: f"{ix('bl_sga')}-{ix('bl_svgp')}", "販管費−サービスの粗利", fmt=F1, bold=False)
t.res("dv_per", "1拠点あたりに必要な利用料の粗利", "万円/年/拠点", lambda: f"{t.a('dv_need', 1)}/{t.a('dv_billed', 1)}", "必要な利用料の粗利÷課金する拠点", fmt=F2, bold=False)
t.res("dv_base", "基盤の粗利（1拠点）", "万円/年/拠点", lambda: f"{A('base_k')}-{A('pc_infra')}", "基盤60−データ基盤12", fmt=F1, bold=False)
t.res("dv_dev", "1台の粗利", "万円/年/台", lambda: f"{A('dev_k')}-{A('perdev')}", "端末24−1台の原価7.2", fmt=F2, bold=False)
t.res("dv_d", "必要な台数（計算）", "台/拠点", lambda: f"({t.a('dv_per', 1)}-{t.a('dv_base', 1)})/{t.a('dv_dev', 1)}", "(1拠点に必要な粗利−基盤の粗利)÷1台の粗利", fmt=F2)
t.res("dv_up", "必要な台数（整数に切り上げ）", "台/拠点", lambda: f"ROUNDUP({t.a('dv_d', 1)},0)", "この台数なら5年目に黒字", fmt=F0)
opd = lambda n: f"{t.a('dv_billed', 1)}*({t.a('dv_base', 1)}+{t.a('dv_dev', 1)}*{n})+{ix('bl_svgp')}-{ix('bl_sga')}"
t.res("dv_op_up", "確認：その台数での5年目の営業利益", "万円/年", lambda: opd(t.a("dv_up", 1)), "課金する拠点×(基盤の粗利＋1台の粗利×台数)＋サービスの粗利−販管費", fmt=F1)
t.res("dv_op_dn", "確認：1台少ない場合の5年目の営業利益", "万円/年", lambda: opd(f"({t.a('dv_up', 1)}-1)"), "同じ式で台数−1", fmt=F1)
t.res("dv_op_15", "確認：15台の場合の5年目の営業利益", "万円/年", lambda: opd(A("check_devs")), "前提9-4の台数。tora_v2.js の1拠点15台（4,301万円）と同じになる", fmt=F1)
t.note("注：台数を増やしても、導入支援・人件費・データ基盤は変わらないとした（tora_v2.js と同じ）。")
t.sec("問い：黒字ラインは結局何台／何拠点？")
r_ans = t.r
bl.ws.merge_cells(start_row=r_ans, start_column=1, end_row=r_ans, end_column=13)
fb = A("first_black10")
ans = lambda: ('="5年目（社員"&TEXT(' + ix("bl_staff") + ',"0")&"人・販管費"&TEXT(' + ix("bl_sga") + ',"#,##0")&"万円）で黒字にするには、通年で課金する拠点が"&TEXT('
               + ix("need_sites") + ',"0.0")&"拠点要る（実際は"&TEXT(' + ix("bl_billed") + ',"0.0")&"）。期末"&TEXT(' + bl.a("dv_sites", 1) + ',"0")&"拠点のまま台数で越えるなら、1拠点"&TEXT('
               + bl.a("dv_up", 1) + ',"0")&"台（計算では"&TEXT(' + bl.a("dv_d", 1) + ',"0.00")&"台）。10台のままなら"&IF(ISNUMBER(' + fb + '),TEXT(' + fb + ',"0")&"年目（期末"&TEXT(INDEX('
               + RNG("bl_sites", here="黒字ライン") + ',1,' + fb + '),"0")&"拠点・社員"&TEXT(INDEX(' + RNG("bl_staff", here="黒字ライン") + ',1,' + fb + '),"0")&"人）に黒字。","10年以内には黒字にならない。")')
put(bl.ws, r_ans, 1, ans, kind="navy", bold=True, wrap=True, size=11, border=False)
bl.ws.row_dimensions[r_ans].height = 48
bl.ws.cell(r_ans, 1).alignment = Alignment(wrap_text=True, vertical="center")
t.r += 1
t.neg_red()


# ======================================================================
# LTV・CAC
# ======================================================================
lc = SS("LTV・CAC", "LTV・CAC：LTVは控えめに7年で約1,571万円、CACは約86万円",
        "LTV＝1拠点の年の粗利×継続年数＋導入支援の粗利。CAC＝(営業の人件費＋広告)÷新規拠点（5年目）。tora_v2.js の unit() と同じ（計算詳細版 p.20）。")
l = lc.a
lc.sec("LTV（1拠点が残す粗利）")
lc.row("lc_gp", "1拠点の年の粗利", lambda: A("gp_site"), "万円/年", "値段と原価（300−84）", fmt=F1)
lc.row("lc_sgp", "導入支援の粗利", lambda: A("setup_gp"), "万円/拠点", "導入支援（200−141）", fmt=F1)
lc.row("lc_years", "計画の継続年数", lambda: A("ltv_plan"), "年", "前提9-1")
lc.row("ltv", "LTV（計画）", lambda: f"{l('lc_gp')}*{l('lc_years')}+{l('lc_sgp')}", "万円/拠点", "216×7年＋59", fmt=F1, bold=True)
lc.row("avg_life", "解約率から見た平均の継続年数", lambda: f"1/{A('churn_used', 1)}", "年", "1÷解約率10%＝10年。7年は控えめ", fmt=F1)
lc.sec("CAC（1拠点を獲得する費用）")
lc.row("lc_year", "計算する年", lambda: A("cac_year"), "年目", "前提9-2", fmt=YR)
lc.row("lc_sn", "営業の人数", lambda: f"INDEX({RNG('h5')},1,{l('lc_year')})", "人", "体制と人件費（5年目3人）")
lc.row("lc_ss", "営業の年収", lambda: f"INDEX({RNG('sal5_y')},1,{l('lc_year')})", "万円/年", "体制と人件費")
lc.row("lc_pay", "営業の人件費", lambda: f"{l('lc_sn')}*{l('lc_ss')}*{A('load')}", "万円/年", "3人×650万円×1.2")
lc.row("lc_ads", "広告・展示会", lambda: f"INDEX({RNG('other_ads')},1,{l('lc_year')})", "万円/年", "前提7-10")
lc.row("lc_new", "新規拠点", lambda: f"INDEX({RNG('new_sites')},1,{l('lc_year')})", "拠点/年", "前提7-1")
lc.row("cac", "CAC", lambda: f"({l('lc_pay')}+{l('lc_ads')})/{l('lc_new')}", "万円/拠点", "(営業3人×650万円×1.2＋広告500)÷新規33", fmt=F2, bold=True)
lc.row("payback", "回収月数", lambda: f"{l('cac')}/({l('lc_gp')}/{A('c_month')})", "か月", "CAC÷(1拠点の年の粗利÷12)", fmt=F2, bold=True)
lc.row("ltv_cac", "LTV÷CAC（計画）", lambda: f"{l('ltv')}/{l('cac')}", "倍", "LTV÷CAC", fmt=F1, bold=True)
lc.pbreak()
lc.sec("継続年数を変えたら")
lc.mhdr(["継続年数", "単位", "LTV（万円/拠点）", "LTV÷CAC（倍）", "式（言葉）", False])
lc.ws.merge_cells(start_row=lc.r - 1, start_column=5, end_row=lc.r - 1, end_column=6)
for i in range(1, 5):
    lc.mrow([((lambda i=i: A("ltv_years", i)), '0"年"', "calc", True), ("万円・倍", None, "gray", False), ((lambda i=i: f"{l('lc_gp')}*{A('ltv_years', i)}+{l('lc_sgp')}"), F0, "calc", False),
             ((lambda i=i: f"({l('lc_gp')}*{A('ltv_years', i)}+{l('lc_sgp')})/{l('cac')}"), F1, "calc", False)],
            "216×継続年数＋59、それをCACで割る", desc_from=5, key=f"ltv_t{i}")
lc.sec("獲得が難しくなったら（CACが増えた場合）")
lc.mhdr(["CACの倍率", "単位", "CAC（万円/拠点）", "LTV÷CAC（倍・計画の年数）", "回収月数（か月）", "式（言葉）"])
for i in range(1, 4):
    lc.mrow([((lambda i=i: A("cac_k", i)), '0"倍"', "calc", True), ("万円・倍・か月", None, "gray", False), ((lambda i=i: f"{l('cac')}*{A('cac_k', i)}"), F0, "calc", False),
             ((lambda i=i: f"{l('ltv')}/({l('cac')}*{A('cac_k', i)})"), F1, "calc", False), ((lambda i=i: f"{l('payback')}*{A('cac_k', i)}"), F1, "calc", False)],
            "CAC×倍率", desc_from=6, key=f"cac_t{i}")
lc.note("注：1〜2年目は創業者が営業するため、CACに創業者の人件費は入れていない（聞かれたらそのまま答える）。")
lc.note("注：SaaS でよく使われる目安は「LTV÷CACが3倍以上、回収12か月以内」。")
lc.neg_red()


# ======================================================================
# 崩れたとき（tora_v2.js の値を貼る）
# ======================================================================
ws = wb["崩れたとき"]
WK = [30, 22, 12, 12, 15, 14, 14, 12, 52]
setup(ws, WK, "崩れたとき：前提を一つずつ崩した結果（tora_v2.js で計算した値）",
      "灰色の数字は tora_v2.js の run() で計算した値を貼ったもの。1年目〜5年目は計算詳細版 p.21 と同じ。10年で見た列は run({years:10}) に同じ変更を入れた値"
      "（新規−30%は6〜10年目の新規と個社設定も0.7倍）。右の列の通りに前提のつまみを変えると、Excel でも同じ結果になる。")
hdr(ws, 3, ["シナリオ", "想定する事態", "黒字化\n（5年で）", "黒字化\n（10年で）", "5年目の営業利益\n（万円）", "最低の現金\n5年で（万円）", "最低の現金\n10年で（万円）", "累計回収\n（10年で）", "Excel で同じ結果を出すには（前提のどこを変えるか）"], height=40)
ws.freeze_panes = "B4"
page(ws)
WHY = ["─", "値引きされた", "定着しなかった", "処理量が想定より多い", "採用市場が厳しい", "営業が想定より遅い", "やってはいけない例", "効果が出た班から追加", "やってはいけない例"]
HOW = ["初期値のまま",
       "8-1 値段の倍率を 0.8",
       "8-2 解約率の上書きに 20%",
       "8-3 原価の倍率を 2",
       "8-4 年収の上乗せを 100",
       "7-1 新規拠点を0.7倍して四捨五入（3,9,18,27,33→2,6,13,19,23）、7-7 個社設定を0.7倍、8-5 採用連動のkに 0.7",
       "7-1 と 7-7 を上と同じく0.7倍（kは空のまま）",
       "2-1 台数を 15（リースの構成は12台＋3台）",
       "8-3 原価の倍率を 60÷7.2（1台 月5万円）",
       "6-14 営業の人数を 0,0,1,3,3（2年目の採用をやめる）"]
names = [s["n"] for s in NV["SC"]] + ["営業の採用を3年目から"]
order = [0, 1, 2, 3, 4, 5, 6, 7, 8]
results = NV["SC"] + [{"n": "営業の採用を3年目から", **NV["salesY3"]}]
WHY = WHY[:5] + ["営業が想定より遅い", "やってはいけない例", "効果が出た班から追加", "やってはいけない例", "営業の立ち上げを遅らせる"]
rr = 4
for i, s_ in enumerate(results):
    o5, o10 = s_["o5"], s_["o10"]
    bad = WHY[i] == "やってはいけない例"
    put(ws, rr, 1, s_["n"], bold=(i == 0 or bad), kind=("red" if bad else "calc"), wrap=True, width=WK[0])
    put(ws, rr, 2, WHY[i], kind=("red" if bad else "gray"), wrap=True, width=WK[1])
    fb5 = o5["firstBlack"]
    fb10 = o10["firstBlack"]
    put(ws, rr, 3, fb5 if fb5 else "5年でなし", YR, "paste", True, align="right")
    put(ws, rr, 4, fb10 if fb10 else "10年でなし", YR, "paste", True, align="right")
    put(ws, rr, 5, o5["op5"], F0, "paste", align="right")
    put(ws, rr, 6, o5["minCash"], F0, "paste", align="right")
    put(ws, rr, 7, o10["minCash"], F0, "paste", align="right")
    put(ws, rr, 8, o10["cumBack"] if o10["cumBack"] else "10年でなし", YR, "paste", align="right")
    put(ws, rr, 9, HOW[i], kind="gray", wrap=True, width=WK[8], size=9)
    rr += 1
REG["sc_first"] = ("崩れたとき", 4, 3)
rr += 1
for t_ in ["注：計画の余裕は小さい。1拠点10台のままでは5年以内に黒字にならず、値引き・解約の増加・AI費用の増加・年収の上昇はどれも赤字を広げ、最低の現金がマイナスになるものもある。",
           "注：特に危ないのは、動画を全部クラウドで解析する使い方と、獲得が遅れているのに計画どおり採用すること。守りは、1拠点の台数を増やすことと、採用を拠点数に合わせて遅らせること。",
           "注：営業の採用を3年目からにすると、2年目の人件費780万円と採用費が浮き、最低の現金は約780万円増える。5年目の営業利益は変わらない（拠点の計画を変えていないため）。",
           "注：黒字化の年は数字（緑）、「なし」は期間内に黒字にならないこと。マイナスは▲（赤）。"]:
    note_row(ws, rr, t_, 9, WK)
    rr += 1
CF_NEG.setdefault("崩れたとき", []).append(f"C4:H{rr}")
ws.conditional_formatting.add(f"C4:D{3 + len(results)}", FormulaRule(formula=["AND(ISNUMBER(C4),C4>0)"], font=Font(color=GREEN, bold=True)))


# ======================================================================
# 資金と黒字倒産
# ======================================================================
cf = TS("資金と黒字倒産", "資金と黒字倒産：自己資金は段階に分けて入れる。黒字でも現金が尽きる場合を先に見る",
        "資金投入は準備期500万円＋1年目2億円＋3年目5,000万円。このモデルは運転資本・税・償却を含まない簡易計算。下の①〜④は、簡易計算の現金からさらに出ていく額を年ごとに見たもの（計算詳細版 p.23）。")
t = cf
t.sec("資金と現金")
t.row("cf_fund", "資金投入", "万円/年", lambda y: A("funding", y), "前提7-16")
t.row("cf_op", "営業利益", "万円/年", lambda y: A("op", y), "損益と現金", fmt=F1)
t.row("cf_cash", "期末の現金（簡易計算）", "万円", lambda y: A("cash", y), "損益と現金", fmt=F1, bold=True)
t.row("cf_line", "最低残高の目安（その年の販管費の3か月分）", "万円", lambda y: f"{A('sga', y)}*{A('min_m')}/{A('c_month')}", "販管費×3÷12", fmt=F1)
t.row("cf_room", "目安との差", "万円", lambda y: f"{t.a('cf_cash', y)}-{t.a('cf_line', y)}", "期末の現金−最低残高の目安", fmt=F1)
t.res("cf_prep", "準備期の資金", "万円", lambda: A("prep_fund"), "試作品・実証。準備期に使い切る")
t.res("cf_total", "資金投入の合計（準備期を含む）", "万円", lambda: f"{t.a('cf_prep', 1)}+SUM({RNG('funding')})", "500＋2億＋5,000万＝2.55億円")
t.res("cf_min", "最低の現金（10年）", "万円", lambda: A("min_cash10"), "損益と現金", fmt=F1)
t.res("cf_line4", "参考：4年目の販管費の3か月分", "万円", lambda: f"INDEX({RNG('sga')},1,4)*{A('min_m')}/{A('c_month')}", "計算詳細版 p.21・23 の最低現金のライン", fmt=F1, bold=False)
t.sec("① 売掛の回収が遅い（月額の後払い・支払いサイト60日）")
t.row("ar", "売掛金の残高（年末）", "万円", lambda y: f"{A('sales', y)}*{A('ar_days')}/{A('c_day')}", "売上×60÷365", fmt=F1)
t.row("ar_inc", "その年に増えた売掛（現金が減る額）", "万円/年", lambda y: f"{t.a('ar', y)}-{t.a('ar', y - 1)}", "今年の残高−前の年の残高", fmt=F1, f1=lambda y: t.a("ar", y))
t.row("cash_ar", "売掛を引いた現金", "万円", lambda y: f"{t.a('cf_cash', y)}-{t.a('ar', y)}", "期末の現金−売掛金の残高", fmt=F1, bold=True)
t.sec("② 採用の先行（拠点が増える前に人件費が出る）")
t.row("pay_next", "翌年の人件費の増加", "万円", lambda y: f"{A('payroll', y + 1)}-{A('payroll', y)}", "翌年の人件費−その年の人件費（10年目は翌年の計画がないので「─」）", years=range(1, 10))
t.row("lead", "先に出る人件費", "万円", lambda y: f"MAX(0,{t.a('pay_next', y)})*{A('hire_lead')}/{A('c_month')}", "翌年の増加×3か月÷12", fmt=F1, years=range(1, 10))
t.row("cash_lead", "先行分を引いた現金", "万円", lambda y: f"{t.a('cf_cash', y)}-{t.a('lead', y)}", "期末の現金−先に出る人件費", fmt=F1, bold=True, years=range(1, 10))
t.sec("③ 端末込みの料金にした場合の端末の先払い（3年目から）")
t.row("dev_new", "新しく買う端末", "台", lambda y: f"IF({t.Y(y)}>={A('dev_start')},{A('new', y)}*{A('devices')},0)", "3年目から 新規拠点×10台")
t.row("dev_pay", "端末の先払い", "万円/年", lambda y: f"{t.a('dev_new', y)}*{A('dev_unit')}", "新しく買う端末×8万円")
unrec = lambda y: "+".join([f"{t.a('dev_pay', y - j)}*MAX(0,1-{j}/{A('dev_recover')})" for j in range(0, 4) if y - j >= 1])
t.row("dev_unrec", "まだ回収していない端末代（年末）", "万円", unrec, "その年に買った分＋1年前の分×(1−1/3)＋2年前の分×(1−2/3)（翌年から3年で均等に回収）", fmt=F1)
t.row("cash_dev", "未回収分を引いた現金", "万円", lambda y: f"{t.a('cf_cash', y)}-{t.a('dev_unrec', y)}", "期末の現金−まだ回収していない端末代", fmt=F1, bold=True)
t.sec("④ 急な拠点増で導入支援の人件費が先に出る（新規が計画より5割多い場合）")
t.row("sg_new", "計画より多い新規拠点", "拠点", lambda y: f"{A('new', y)}*{A('surge')}", "新規×50%", fmt=F1)
t.row("sg_cost", "先に出る導入支援の原価", "万円", lambda y: f"{t.a('sg_new', y)}*{A('setup_cost')}", "多い新規×141万円", fmt=F1)
t.row("sg_dep", "着手金で先に入る額", "万円", lambda y: f"{t.a('sg_new', y)}*{A('setup_fee')}*{A('deposit')}", "多い新規×200万円×50%", fmt=F1)
t.row("sg_gap", "差し引き先に出る現金", "万円", lambda y: f"{t.a('sg_cost', y)}-{t.a('sg_dep', y)}", "原価−着手金", fmt=F1)
t.row("sg_cs", "追加で要る導入コンサルタント", "人", lambda y: f"{t.a('sg_new', y)}/{A('cs_cap')}", "多い新規÷6（採用はさらに先行する）", fmt=F1)
t.row("cash_sg", "差し引きを引いた現金", "万円", lambda y: f"{t.a('cf_cash', y)}-{t.a('sg_gap', y)}", "期末の現金−差し引き先に出る現金", fmt=F1, bold=True)
t.sec("①〜④のいちばん低い現金")
t.res("min_ar", "① 売掛を引いた現金の最低", "万円", lambda: f"MIN({RNG('cash_ar', here=cf.name)})", "", fmt=F1)
t.res("min_lead", "② 先行分を引いた現金の最低", "万円", lambda: f"MIN({RNG('cash_lead', 9, here=cf.name)})", "1〜9年目", fmt=F1)
t.res("min_dev", "③ 未回収分を引いた現金の最低", "万円", lambda: f"MIN({RNG('cash_dev', here=cf.name)})", "", fmt=F1)
t.res("min_sg", "④ 差し引きを引いた現金の最低", "万円", lambda: f"MIN({RNG('cash_sg', here=cf.name)})", "", fmt=F1)
for k_, lab_, n_ in [("cash_ar", "①", 10), ("cash_lead", "②", 9), ("cash_dev", "③", 10), ("cash_sg", "④", 10)]:
    t.res(f"below_{k_}", f"{lab_} 最低残高の目安（販管費3か月分）を下回る年の数", "年",
          (lambda k_=k_, n_=n_: f"SUMPRODUCT(--({RNG(k_, n_, here=cf.name)}<{RNG('cf_line', n_, here=cf.name)}))"), "その年の販管費×3÷12 を下回る年を数える", fmt=F0, bold=False)
t.sec("黒字倒産が起こる場合と備え")
for t_ in [
    "① 売掛の回収が遅い：月額の後払いで支払いサイト60日なら、売上×60÷365の運転資金が要る（5年目で約4,400万円）。備え：利用料は年額前払い。導入支援は着手金50%。",
    "② 採用の先行：導入コンサルタントや営業を、拠点が増える前に採ると人件費が先に出る。備え：採用は拠点数の実績に合わせて遅らせる（採用連動。前提8-5）。期末で社員1人約4拠点の見込みが立ってから採る。",
    "③ 端末込みの料金：当社が端末をまとめて買うと、1台8万円×新規の台数を先に払い、料金で3年かけて回収する。備え：本契約はお客様名義のリースを基本にし、端末込みは現金に余裕がある年だけ選べるようにする。",
    "④ 急な拠点増：新規が計画より多いと、導入支援の人件費と交通費が売上より先に出て、導入コンサルタントの採用も先に要る。備え：着手金50%。新規の受け入れは導入コンサルタントの人数×6を上限にする。",
    "⑤ 税金・賞与の支払い月：賞与は年2回、法人税等は決算から2か月以内にまとめて出る。黒字化した後は中間納付もある（このモデルは税を含まない）。備え：資金繰り表で12か月先までの入出金を管理する。",
    "共通の備え：自己資金は段階に分けて入れる（準備期500万円→実証に合格したら1年目2億円→3年目に継続率と台数の線を越えていれば5,000万円）。最低残高は販管費の3か月分を下回らないようにし、下回りそうなら先に採用を止める。",
    "注：このモデルは運転資本（売掛・前受）・税・償却を含まない簡易計算。①〜④は互いに重なりうる。",
]:
    t.note(t_, kind=("calc" if t_.startswith(("①", "②", "③", "④", "⑤", "共通")) else "gray"))
t.neg_red()


# ======================================================================
# 市場規模
# ======================================================================
ws = wb["市場規模"]
WM = [31, 17, 13, 16, 15, 15, 13, 11, 46]
HM = "市場規模"
setup(ws, WM, "市場規模：ターゲット市場（電気工事・電気通信工事の中堅以上）約141億円",
      "業種×規模で、顧客数×1社の年額×年1回を足し上げる。tora_market.js の SEG・SOM・estimate() と同じ（計算詳細版 p.24・25）。")
hdr(ws, 3, ["業種", "規模", "顧客数（社）", "1社の年額（万円/年）", "ターゲット（1＝入れる）", "市場（万円/年）", "市場（億円/年）", "確度", "メモ・式（言葉）"], height=34)
ws.freeze_panes = "A4"
page(ws)
mr = 4
for i, seg in enumerate(SEG):
    ind, size = seg[0], seg[1]
    small_e = ind == "電気工事会社" and size.startswith("小")
    put(ws, mr, 1, ind)
    put(ws, mr, 2, size)
    put(ws, mr, 3, (lambda i=i: "=" + A(f"seg{i}", 1)), F0, align="right")
    put(ws, mr, 4, (lambda i=i: "=" + A(f"seg{i}", 2)), F0, align="right")
    put(ws, mr, 5, (lambda i=i: "=" + A(f"seg{i}", 3)), F0, align="center")
    put(ws, mr, 6, (lambda mr=mr: f"=C{mr}*D{mr}"), F0, align="right")
    put(ws, mr, 7, (lambda mr=mr: f"=F{mr}/{A('c_oku')}"), F1, align="right")
    put(ws, mr, 8, "統計−仮定" if small_e else "仮定", align="center")
    put(ws, mr, 9, seg[6] + ("（65,497社−大300−中2,100）" if small_e else "") + "。顧客数×1社の年額", kind="gray", wrap=True, width=WM[8], size=9)
    reg(f"mk{i}", HM, mr, 3)
    mr += 1
seg_first, seg_last = 4, mr - 1
rows_m = [
    ("tam", "全セルの合計（参考）", lambda: f"=SUM(F{seg_first}:F{seg_last})", lambda: f"=SUM(C{seg_first}:C{seg_last})", "8つのセルの合計（電線工場と小規模を含む）"),
    ("sam", "ターゲット市場（ターゲット＝1のセル）", lambda: f"=SUMPRODUCT(E{seg_first}:E{seg_last},F{seg_first}:F{seg_last})", lambda: f"=SUMPRODUCT(E{seg_first}:E{seg_last},C{seg_first}:C{seg_last})",
     "電気工事・電気通信工事の大と中。顧客数×年額をターゲットのセルだけ足す"),
]
for key, label, f6, f3, desc in rows_m:
    put(ws, mr, 1, label, bold=True)
    ws.merge_cells(start_row=mr, start_column=1, end_row=mr, end_column=2)
    put(ws, mr, 3, f3, F0, bold=True, align="right")
    put(ws, mr, 6, f6, F0, bold=True, align="right")
    put(ws, mr, 7, (lambda mr=mr: f"=F{mr}/{A('c_oku')}"), F1, bold=True, align="right")
    put(ws, mr, 9, desc, kind="gray", wrap=True, width=WM[8], size=9)
    reg(key, HM, mr, 6)
    reg(key + "_n", HM, mr, 3)
    reg(key + "_oku", HM, mr, 7)
    mr += 1
mr += 1
sec_row(ws, mr, "SOM（1〜5年目に狙うセルの漏斗）", 9)
mr += 1
hdr(ws, mr, ["セル", "顧客数（社）", "地域の割合", "会える率", "成約率", "1社の拠点数", "獲得する拠点", "売上（万円/年）", "式（言葉）"], height=34)
mr += 1
som_first = mr
for j, (ci, region, reach, close, per, note) in enumerate(SOM):
    put(ws, mr, 1, f"{SEG[ci][0]}・{SEG[ci][1]}", wrap=True, width=WM[0])
    put(ws, mr, 2, (lambda ci=ci: "=" + A(f"mk{ci}", 1, here=HM)), F0, align="right")
    for k in range(4):
        put(ws, mr, 3 + k, (lambda j=j, k=k: "=" + A(f"som{j}", k + 1)), F2, align="right")
    put(ws, mr, 7, (lambda mr=mr: f"=B{mr}*C{mr}*D{mr}*E{mr}*F{mr}"), F2, align="right")
    put(ws, mr, 8, (lambda mr=mr: f"=G{mr}*{A('price_site0')}"), F0, align="right")
    put(ws, mr, 9, "顧客数×地域×会える率×成約率×1社の拠点数。売上＝拠点×300万円。" + note, kind="gray", wrap=True, width=WM[8], size=9)
    mr += 1
put(ws, mr, 1, "SOM の合計", bold=True)
put(ws, mr, 7, (lambda mr=mr: f"=SUM(G{som_first}:G{mr - 1})"), F2, bold=True, align="right")
put(ws, mr, 8, (lambda mr=mr: f"=SUM(H{som_first}:H{mr - 1})"), F0, bold=True, align="right")
put(ws, mr, 9, "獲得する拠点の合計と売上（約88拠点・約2.64億円）", kind="gray", wrap=True, width=WM[8], size=9)
reg("som_sites", HM, mr, 7)
reg("som_value", HM, mr, 8)
mr += 1
put(ws, mr, 1, "5年目の売上（期末の拠点×300万円）÷ターゲット市場", bold=False, wrap=True, width=WM[0])
ws.merge_cells(start_row=mr, start_column=1, end_row=mr, end_column=2)
put(ws, mr, 7, (lambda: f"={A('close', 5)}*{A('price_site0')}/{A('sam', 1, here=HM)}"), P2, bold=True, align="right")
put(ws, mr, 9, "81拠点×300万円÷141億円", kind="gray", size=9)
mr += 2
ws.row_breaks.append(Break(id=mr - 1))
sec_row(ws, mr, "参考：小規模向けのライト版（基盤＋3台）で一部を取れた場合の上振れ（計画に入れない）", 9)
mr += 1
hdr(ws, mr, ["取れる割合", "小規模の電気工事会社（社）", "取れる社数（社）", "ライト版の年額（万円/年）", "売上（万円/年）", "売上（億円/年）", False, False, "式（言葉）"], height=34)
mr += 1
for i in range(1, 4):
    put(ws, mr, 1, (lambda i=i: "=" + A("light_share", i)), P1, bold=True)
    put(ws, mr, 2, (lambda: "=" + A("mk5", 1, here=HM)), F0, align="right")
    put(ws, mr, 3, (lambda mr=mr: f"=A{mr}*B{mr}"), F0, align="right")
    put(ws, mr, 4, (lambda: f"={A('base_fee')}+{A('dev_fee')}*{A('light_devs')}"), F0, align="right")
    put(ws, mr, 5, (lambda mr=mr: f"=C{mr}*D{mr}"), F0, align="right")
    put(ws, mr, 6, (lambda mr=mr: f"=E{mr}/{A('c_oku')}"), F1, align="right")
    put(ws, mr, 9, "小規模の社数×取れる割合×(基盤60＋端末24×3台)", kind="gray", size=9)
    mr += 1
mr += 1
for t_ in ["注：65,497社は国の統計（建設業許可業者数調査）。大・中・小の内訳は仮定で、業界団体の名簿と経済センサスで数え直す。",
           "注：小規模（30人未満）を外すのは、書類が少ないからではなく、1拠点10台が要らないため。全体の表の小規模の年額60万円は tora_market.js の置き方（さらに小さい構成）。",
           "注：ライト版は計画に入れない上振れ。取れる割合は仮定。"]:
    note_row(ws, mr, t_, 9, WM)
    mr += 1
CF_NEG.setdefault(HM, []).append(f"B4:H{mr}")


# ======================================================================
# 照合
# ======================================================================
ws = wb["照合"]
WC = [14, 40, 8, 11, 16, 16, 13, 9, 30]
HC = "照合"
setup(ws, WC, "照合：Excel の結果と tora_v2.js の値（node で計算して貼った値）",
      "E列＝このブックの式の結果、F列＝tora_v2.js の値（灰色・貼り付け）。差の絶対値が前提12-1以下なら「一致」。前提のつまみ（8-1〜8-6）が初期値のときに一致する。つまみを動かすと「不一致」になるのが正しい。")
hdr(ws, 3, ["区分", "項目", "年", "単位", "Excel の値", "tora_v2.js の値", "差（Excel−JS）", "判定", "Excel のセル"], height=30)
ws.freeze_panes = "A4"
page(ws)
R10 = NV["run10"]["rows"]
U = NV["unit"]
CHECK = [
    ("時給", "職長の1時間の人件費", None, "円/時間", "j_h", 1, NV["H"]["electrician"] * 10000),
    ("時給", "施工管理者の1時間の人件費", None, "円/時間", "j_h", 2, NV["H"]["manager"] * 10000),
    ("お客様", "1台1日の効果", None, "万円/台/日", "per_day_man", 1, NV["cv"]["perDeviceDay"]),
    ("お客様", "お客様の効果（年）", None, "万円/年", "docs", 1, NV["cv"]["docs"]),
    ("お客様", "浮く時間", None, "時間/年", "freed_h", 1, NV["cv"]["freedHours"]),
    ("お客様", "お客様の支払い", None, "万円/年", "cost", 1, NV["cv"]["cost"]),
    ("お客様", "リースの合計", None, "万円/年", "pay_lease", 1, NV["leaseSite"]),
    ("お客様", "倍率", None, "倍", "ratio", 1, NV["cv"]["ratio"]),
    ("お客様", "元が取れる分", None, "分/台/日", "be", 1, NV["cv"]["breakEvenMin"]),
    ("お客様", "手戻りの削減（参考）", None, "万円/年", "rw", 1, NV["cv"]["rework"]),
    ("値段と原価", "1拠点の利用料", None, "万円/年", "price_site", 1, NV["priceSite"]),
    ("値段と原価", "1台の原価（年）", None, "万円/年/台", "perdev", 1, NV["cloudPerDevice"]),
    ("値段と原価", "1拠点の原価", None, "万円/年", "cogs_site", 1, NV["cogsSite"]),
    ("値段と原価", "1拠点の粗利", None, "万円/年", "gp_site", 1, U["gpYear"]),
    ("値段と原価", "粗利率（利用料）", None, "割合", "margin_site", 1, NV["marginSite"]),
]
for i, m in enumerate(NV["margins"], 1):
    CHECK.append(("値段と原価", f"粗利率：{m['n']}台・大事な場面だけ", None, "割合", f"mg{i}", 2, m["part"]))
    CHECK.append(("値段と原価", f"粗利率：{m['n']}台・全編を毎秒解析", None, "割合", f"mg{i}", 3, m["full"]))
CHECK += [
    ("導入支援", "1人日の人件費", None, "万円/人日", "csday", 1, NV["CSDAY"]),
    ("導入支援", "導入支援の原価", None, "万円/拠点", "setup_cost", 1, NV["setupCost"]),
    ("導入支援", "導入支援の粗利", None, "万円/拠点", "setup_gp", 1, U["setupGp"]),
]
METRICS = [("拠点", "期末の拠点", "拠点", "close", "sites"), ("売上", "売上", "万円", "sales", "sales"), ("粗利", "粗利", "万円", "gp", "gp"),
           ("販管費", "販管費", "万円", "sga", "sga"), ("営業利益", "営業利益", "万円", "op", "op"), ("累計", "累計の営業損益", "万円", "cum", "cum"),
           ("現金", "期末の現金", "万円", "cash", "cash")]
for cat, label, unit, key, nk in METRICS:
    for y in range(1, 11):
        CHECK.append((cat, label, y, unit, key, y, R10[y - 1][nk]))
for y in range(1, 11):
    CHECK.append(("体制", "社員数", y, "人", "staff", y, NV["team"][y - 1]["staff"]))
for y in range(1, 11):
    CHECK.append(("体制", "人件費", y, "万円", "payroll", y, NV["team"][y - 1]["payroll"]))
for y in range(1, 11):
    CHECK.append(("体制", "採用費", y, "万円", "recruit", y, NV["team"][y - 1]["recruit"]))
CHECK += [
    ("結果", "黒字化の年（10年で）", None, "年目", "first_black10", 1, NV["run10"]["firstBlack"]),
    ("結果", "累計回収の年（10年で）", None, "年目", "cum_back10", 1, NV["run10"]["cumBack"]),
    ("結果", "最低の現金（10年）", None, "万円", "min_cash10", 1, NV["run10"]["minCash"]),
    ("結果", "累計赤字の底", None, "万円", "cum_bottom", 1, NV["run10"]["cumPeak"]),
    ("LTV・CAC", "LTV（7年）", None, "万円", "ltv", 1, NV["ltv"][2]),
]
for i, yy in enumerate([3, 5, 7, 10], 1):
    CHECK.append(("LTV・CAC", f"LTV（{yy}年）", None, "万円", f"ltv_t{i}", 1, NV["ltv"][i - 1]))
CHECK += [
    ("LTV・CAC", "CAC", None, "万円", "cac", 1, U["cac"]),
    ("LTV・CAC", "回収月数", None, "か月", "payback", 1, U["paybackMonths"]),
    ("市場規模", "ターゲット市場", None, "万円/年", "sam", 1, NV["market"]["sam"]),
    ("市場規模", "全セルの合計", None, "万円/年", "tam", 1, NV["market"]["tamSeg"]),
    ("市場規模", "ターゲットの社数", None, "社", "sam_n", 1, NV["market"]["samSites"]),
    ("市場規模", "SOM の拠点", None, "拠点", "som_sites", 1, NV["market"]["somSites"]),
    ("市場規模", "SOM の売上", None, "万円/年", "som_value", 1, NV["market"]["somValue"]),
]
# 黒字化の年（5年）は tora_v2.js では null（5年で黒字にならない）
cr = 5
put(ws, 4, 1, "合計", bold=True)
put(ws, 4, 2, (lambda: f'="照合 "&COUNTA(H{5}:H{4 + len(CHECK) + 1})&" 項目：一致 "&COUNTIF(H5:H{4 + len(CHECK) + 1},"一致")&"・不一致 "&COUNTIF(H5:H{4 + len(CHECK) + 1},"不一致")'), bold=True, kind="navy")
ws.merge_cells("B4:G4")
put(ws, 4, 8, (lambda: f'=IF(COUNTIF(H5:H{4 + len(CHECK) + 1},"不一致")=0,"すべて一致","不一致あり")'), bold=True, kind="green", align="center")
REG["chk_summary"] = (HC, 4, 8)
for cat, label, y, unit, key, yi, nv in CHECK:
    put(ws, cr, 1, cat, kind="gray")
    put(ws, cr, 2, label)
    put(ws, cr, 3, y, YR if y else None, align="center")
    put(ws, cr, 4, unit, kind="gray")
    sheet, row, col = REG[key]
    fmt = P2 if unit == "割合" else (F4 if (isinstance(nv, float) and abs(nv) < 1000 and abs(nv - round(nv)) > 1e-9) else F2)
    if unit == "年目":
        fmt = YR
    put(ws, cr, 5, (lambda key=key, yi=yi: "=" + A(key, yi, here=HC)), fmt, align="right")
    put(ws, cr, 6, nv, fmt, "paste", align="right")
    put(ws, cr, 7, (lambda cr=cr: f'=IF(AND(ISNUMBER(E{cr}),ISNUMBER(F{cr})),E{cr}-F{cr},"─")'), '0.0E+00;-0.0E+00;0', align="right")
    put(ws, cr, 8, (lambda cr=cr: f'=IF(ISNUMBER(G{cr}),IF(ABS(G{cr})<={A("tol")},"一致","不一致"),IF(E{cr}=F{cr},"一致","不一致"))'), align="center", bold=True)
    put(ws, cr, 9, f"{sheet} {CL(col + yi - 1)}{row}", kind="gray", size=9)
    cr += 1
put(ws, cr, 1, "結果", kind="gray")
put(ws, cr, 2, "黒字化の年（5年で）")
put(ws, cr, 4, "年目", kind="gray")
put(ws, cr, 5, (lambda: "=" + A("first_black5", 1, here=HC)), YR, align="right")
put(ws, cr, 6, "5年でなし", None, "paste", align="right")
put(ws, cr, 7, (lambda cr=cr: f'=IF(AND(ISNUMBER(E{cr}),ISNUMBER(F{cr})),E{cr}-F{cr},"─")'), align="right")
put(ws, cr, 8, (lambda cr=cr: f'=IF(ISNUMBER(G{cr}),IF(ABS(G{cr})<={A("tol")},"一致","不一致"),IF(E{cr}=F{cr},"一致","不一致"))'), align="center", bold=True)
put(ws, cr, 9, "損益と現金（tora_v2.js の run() は firstBlack＝null）", kind="gray", size=9)
cr += 2
note_row(ws, cr, "注：tora_v2.js の値は deck フォルダで node を実行して出した（V.H、V.customerValue()、V.unit()、V.team(10)、V.cogsSite()、V.marginSite()、V.run({years:10})、tora_market.js の estimate()）。差は Excel−JS。", 9, WC)
ws.conditional_formatting.add(f"H5:H{cr}", CellIsRule(operator="equal", formula=['"不一致"'], font=Font(color=RED, bold=True)))
ws.conditional_formatting.add(f"H5:H{cr}", CellIsRule(operator="equal", formula=['"一致"'], font=Font(color=GREEN, bold=True)))
CF_NEG.setdefault(HC, []).append(f"E5:F{cr}")


# ======================================================================
# 使い方
# ======================================================================
ws = wb["使い方"]
WU = [22, 70, 26, 46]
setup(ws, WU, "使い方：カンツギ 事業計画書「計算詳細版」の数字の裏取り",
      "このブックは、計算詳細版の数字が計算モデル（deck/tora_v2.js）と同じ式で出ることを、すべての計算過程が追える形で示すためのもの。入力は「前提」だけにあり、他のシートはすべて式。最後の「照合」で、Excel の結果と tora_v2.js の値が一致することを確かめている。")
page(ws, title_rows=None)
ur = 4
sec_row(ws, ur, "シートの一覧", 4)
ur += 1
hdr(ws, ur, ["シート", "中身", "計算詳細版のページ", "tora_v2.js の対応"])
ur += 1
LIST = [
    ("前提", "入力をすべて集めた（青字）。確度と出典を付けた。年ごとの項目は1〜10年目", "p.4・16・17・32", "WAGE・PLAN・SAVE・CLOUD・SETUP_TASKS・ROLES・EXT・GROWTH・LOAD・RECRUIT_RATE・ADVISORS"),
    ("時給", "職長・施工管理者の1時間の人件費＝年収×会社負担÷年間総実労働時間", "p.5", "H（hourly）"),
    ("お客様の効果", "1台1日の効果、年の効果、支払い、倍率、元が取れる分、感度表、手戻り（参考）", "p.6・7・8", "customerValue()"),
    ("値段と原価", "1拠点の売上、1台の原価の内訳と検算、粗利率、広い原価の粗利率、台数と使い方、追加料金の案", "p.9・10・11・12", "priceSite・CLOUD・cloudPerDevice・cogsSite・marginSite"),
    ("導入支援", "6つの作業×人日、日額、原価・値段・粗利、お客様の1年目の支払い", "p.13", "SETUP_TASKS・setupDays・CSDAY・setupTravel"),
    ("体制と人件費", "8つの役割の年収×人数、人件費、採用費、つまみの反映、導入コンサルタントの上限", "p.17", "ROLES・EXT.heads・team()"),
    ("拠点と売上", "期首・解約・新規・期末、既存と新規の利用料、導入支援、個社設定、売上", "p.14・15", "run() の拠点と売上"),
    ("損益と現金", "原価、粗利、販管費、営業利益、累計、現金。黒字化の年・累計回収の年・最低の現金", "p.3・18・30", "run() の損益と現金"),
    ("黒字ライン", "損益分岐点売上、社員1人あたりの拠点、黒字に必要な拠点、5年目に要る台数", "p.19・22", "run() の bep・sitesPerStaff"),
    ("LTV・CAC", "LTV（3・5・7・10年）、CAC、回収月数", "p.20", "unit()"),
    ("崩れたとき", "前提を一つずつ崩した結果（tora_v2.js の値を貼った）と、Excel で再現する方法", "p.21・22", "run(o)・shortfall()"),
    ("資金と黒字倒産", "資金投入と現金、黒字でも現金が尽きる場合（売掛・採用の先行・端末の先払い・急な拠点増・税と賞与）と備え", "p.23", "run() の cash"),
    ("市場規模", "業種×規模の表、ターゲット市場 約141億円、SOM の漏斗、ライト版の上振れ（参考）", "p.24・25", "tora_market.js の SEG・SOM・estimate()"),
    ("照合", "Excel の主な結果と tora_v2.js の値を並べ、差がゼロ（丸め誤差）であることを示す", "─", "node で出した値"),
]
for row_ in LIST:
    for j, v in enumerate(row_):
        put(ws, ur, 1 + j, v, bold=(j == 0), wrap=True, width=WU[j], kind=("calc" if j < 2 else "gray"))
    ur += 1
ur += 1
sec_row(ws, ur, "色の見方", 4)
ur += 1
LEG = [("青字", "入力。前提シートだけにある。ここを書き換えると全部のシートが計算し直される", "input"),
       ("黒字", "計算（式）", "calc"),
       ("赤字（▲）", "マイナスの数字（赤字の金額など）", "red"),
       ("緑字", "黒字化の年。損益と現金の営業利益の行では、初めて黒字になる年を緑で示す", "green"),
       ("灰色の数字", "tora_v2.js を node で計算して貼った値（崩れたとき・照合の数字の列）。式ではない", "paste"),
       ("紺の行", "見出し", "navy")]
for a_, b_, k_ in LEG:
    c_ = put(ws, ur, 1, a_, kind=k_, bold=True)
    if k_ == "navy":
        c_.fill = HDR_FILL
        c_.font = font("FFFFFF", True)
    put(ws, ur, 2, b_, wrap=True, width=WU[1])
    ur += 1
ur += 1
sec_row(ws, ur, "確度の見方（前提シートの「確度」の列）", 4)
ur += 1
for a_, b_ in [("統計", "国の統計（賃金構造基本統計調査、毎月勤労統計調査、建設業許可業者数調査など）"),
               ("制度", "法律や保険料率など、制度で決まる値（社会保険料の会社負担、残業の割増）"),
               ("公表", "メーカー・クラウド・AIの公表価格"),
               ("見積", "社内の見積（値段、台数、人数、年収、経費など）"),
               ("仮説", "実証で測るもの（削減時間、手戻りの削減、会える率・成約率など）"),
               ("仮定", "計算のために置いた値（為替、シナリオのつまみ、参考の表の軸など）"),
               ("換算", "単位の換算（60分＝1時間など）。前提ではないが、式の中に数字を書かないために置いた")]:
    put(ws, ur, 1, a_, bold=True)
    put(ws, ur, 2, b_, wrap=True, width=WU[1])
    ur += 1
ur += 1
sec_row(ws, ur, "決まりと注意", 4)
ur += 1
for t_ in ["・金額は、ことわりがなければ万円。各行に単位を書き、右の列に式（言葉）を書いた。",
           "・前提のつまみ（8-1〜8-6）が初期値のとき、tora_v2.js の計画（run({years:10})）と一致する。5年目までの値は run()（5年）と同じ。",
           "・つまみを変えれば、崩れたときのシナリオを Excel で再現できる（崩れたときシートの右の列）。",
           "・このモデルは運転資本・税・償却を含まない簡易計算。",
           "・元のファイル：deck/tora_v2.js（計算モデル）、deck/tora_market.js（市場規模）、deck/build_tora_detail.js（計算詳細版の作り）。モデルと資料のファイルは書き換えていない。",
           "・フォントは BIZ UDPゴシック を指定している（無い環境では別のフォントで表示される）。PDF は BIZ UDPゴシック の無い環境で書き出したので IPA Pゴシック で表示している。"]:
    note_row(ws, ur, t_, 4, WU, kind="calc")
    ur += 1


# ======================================================================
# 仕上げ：保留していた式を入れる、赤字の色、印刷
# ======================================================================
for cell, fn in PENDING:
    cell.value = fn()

for sheet, ranges in CF_NEG.items():
    for rg in ranges:
        wb[sheet].conditional_formatting.add(rg, CellIsRule(operator="lessThan", formula=["0"], font=Font(color=RED)))

# 黒字化の年のセルを緑の太字に（数字のとき）
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT, "formulas pending:", len(PENDING))
json.dump(REG, open(WORK + "/reg.json", "w", encoding="utf-8"), ensure_ascii=False)
