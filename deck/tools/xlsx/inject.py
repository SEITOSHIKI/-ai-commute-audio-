# openpyxl で作ったブック（式・書式はそのまま）に、LibreOffice で再計算した値を「計算済みの値」として入れる
import sys, zipfile, posixpath
from lxml import etree
from openpyxl import load_workbook
src, recalced, dst = sys.argv[1], sys.argv[2], sys.argv[3]
NS = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
RNS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
vals = load_workbook(recalced, data_only=True)
zi = zipfile.ZipFile(src)
wbx = etree.fromstring(zi.read("xl/workbook.xml"))
rels = etree.fromstring(zi.read("xl/_rels/workbook.xml.rels"))
rid2target = {r.get("Id"): r.get("Target") for r in rels}
part2sheet = {}
for sh in wbx.find(f"{{{NS}}}sheets"):
    tgt = rid2target[sh.get(f"{{{RNS}}}id")]
    part2sheet[posixpath.normpath(posixpath.join("xl", tgt)) if not tgt.startswith("/") else tgt.lstrip("/")] = sh.get("name")
n = 0; missing = 0
with zipfile.ZipFile(dst, "w", zipfile.ZIP_DEFLATED) as zo:
    for item in zi.infolist():
        data = zi.read(item.filename)
        if item.filename in part2sheet:
            ws = vals[part2sheet[item.filename]]
            root = etree.fromstring(data)
            for c in root.iter(f"{{{NS}}}c"):
                f = c.find(f"{{{NS}}}f")
                if f is None:
                    continue
                v = c.find(f"{{{NS}}}v")
                if v is None:
                    v = etree.SubElement(c, f"{{{NS}}}v")
                val = ws[c.get("r")].value
                if isinstance(val, bool):
                    c.set("t", "b"); v.text = "1" if val else "0"
                elif isinstance(val, (int, float)):
                    if "t" in c.attrib: del c.attrib["t"]
                    v.text = repr(float(val)) if isinstance(val, float) else str(val)
                elif val is None:
                    c.set("t", "str"); v.text = ""; missing += 1
                else:
                    s = str(val)
                    if s.startswith("#"):
                        c.set("t", "e")
                    else:
                        c.set("t", "str")
                    v.text = s
                n += 1
            data = etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)
        zo.writestr(item, data)
print("cached values written:", n, "| empty results:", missing)
